"""Sinh animation thật của PowerPoint (p:timing) và hiệu ứng chuyển trang (p:transition, kể cả Morph).

Mỗi slide nhận danh sách `steps`: mỗi bước là list (shape, hiệu_ứng). Các shape trong cùng
một bước xuất hiện cùng lúc; các bước nối tiếp nhau.
  mode = "auto"  : tự chạy khi vào slide (After Previous)
  mode = "click" : tiêu đề tự chạy, mỗi bước sau chờ một cú nhấp
"""
from __future__ import annotations

from lxml import etree
from pptx.oxml.ns import qn

P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
NSDECL = f'xmlns:p="{P_NS}" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'

# presetID PowerPoint chuẩn: Fade=10, Float In=42, Wipe=22, Zoom=53, Fly In=2
PRESETS = {"fade": (10, 0), "float": (42, 0), "wipe": (22, 4), "zoom": (53, 16), "fly": (2, 4)}


class _Ids:
    def __init__(self):
        self.n = 2

    def __call__(self):
        self.n += 1
        return self.n


def _set_vis(ids, spid):
    return (f'<p:set><p:cBhvr><p:cTn id="{ids()}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/>'
            f'</p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst>'
            f'<p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/>'
            f'</p:to></p:set>')


def _anim(ids, spid, attr, v0, v1, dur):
    return (f'<p:anim calcmode="lin" valueType="num"><p:cBhvr additive="base"><p:cTn id="{ids()}" dur="{dur}" '
            f'fill="hold"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>{attr}'
            f'</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav tm="0"><p:val><p:strVal val="{v0}"/>'
            f'</p:val></p:tav><p:tav tm="100000"><p:val><p:strVal val="{v1}"/></p:val></p:tav></p:tavLst></p:anim>')


def _filter(ids, spid, flt, dur):
    return (f'<p:animEffect transition="in" filter="{flt}"><p:cBhvr><p:cTn id="{ids()}" dur="{dur}"/>'
            f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>')


def _effect(ids, spid, eff, node, dur):
    pid, sub = PRESETS.get(eff, PRESETS["fade"])
    body = _set_vis(ids, spid)
    if eff == "float":
        body += _filter(ids, spid, "fade", dur) + _anim(ids, spid, "ppt_x", "#ppt_x", "#ppt_x", dur) \
            + _anim(ids, spid, "ppt_y", "#ppt_y+.1", "#ppt_y", dur)
        extra = ' decel="100000"'
    elif eff == "wipe":
        body += _filter(ids, spid, "wipe(up)", dur)
        extra = ""
    elif eff == "zoom":
        body += _anim(ids, spid, "ppt_w", "0", "#ppt_w", dur) + _anim(ids, spid, "ppt_h", "0", "#ppt_h", dur) \
            + _filter(ids, spid, "fade", dur)
        extra = ""
    elif eff == "fly":
        body += _anim(ids, spid, "ppt_x", "#ppt_x", "#ppt_x", dur) \
            + _anim(ids, spid, "ppt_y", "1+#ppt_h/2", "#ppt_y", dur)
        extra = ' decel="100000"'
    else:
        body += _filter(ids, spid, "fade", dur)
        extra = ""
    return (f'<p:par><p:cTn id="{ids()}" presetID="{pid}" presetClass="entr" presetSubtype="{sub}" fill="hold" '
            f'grpId="0" nodeType="{node}"{extra}><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'{body}</p:childTnLst></p:cTn></p:par>')


def _group(ids, delay, inner):
    return (f'<p:par><p:cTn id="{ids()}" fill="hold"><p:stCondLst><p:cond delay="{delay}"/></p:stCondLst>'
            f'<p:childTnLst>{inner}</p:childTnLst></p:cTn></p:par>')


def add_animations(slide, steps, mode="auto", dur=550, gap=120):
    """steps: list[list[(shape, effect)]]."""
    steps = [s for s in steps if s]
    if not steps or mode == "none":
        return
    ids = _Ids()
    clicks: list[list[list]] = []  # mỗi phần tử = một lần nhấp, gồm nhiều bước nối tiếp
    if mode == "click":
        clicks.append([steps[0]])
        clicks += [[s] for s in steps[1:]]
    else:
        clicks.append(steps)

    pars = []
    text_spids = []
    for ci, chain in enumerate(clicks):
        t = 0
        groups = []
        for si, step in enumerate(chain):
            effs = []
            for k, (shp, eff) in enumerate(step):
                if ci == 0 and si == 0 and k == 0:
                    node = "afterEffect"
                elif k == 0:
                    node = "clickEffect" if (mode == "click" and si == 0) else "afterEffect"
                else:
                    node = "withEffect"
                effs.append(_effect(ids, shp.shape_id, eff, node, dur))
                if getattr(shp, "has_text_frame", False) and shp.has_text_frame and shp.text_frame.text.strip():
                    text_spids.append((shp.shape_id, _has_fill(shp)))
            groups.append(_group(ids, t, "".join(effs)))
            t += dur + gap
        cond = ('<p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond>'
                if ci == 0 else '<p:cond delay="indefinite"/>')
        pars.append(f'<p:par><p:cTn id="{ids()}" fill="hold"><p:stCondLst>{cond}</p:stCondLst>'
                    f'<p:childTnLst>{"".join(groups)}</p:childTnLst></p:cTn></p:par>')

    bld = "".join(f'<p:bldP spid="{sid}" grpId="0"{" animBg=\"1\"" if bg else ""}/>' for sid, bg in text_spids)
    xml = (f'<p:timing {NSDECL}><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">'
           f'<p:childTnLst><p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq">'
           f'<p:childTnLst>{"".join(pars)}</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0">'
           f'<p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0">'
           f'<p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq></p:childTnLst></p:cTn></p:par>'
           f'</p:tnLst>{"<p:bldLst>" + bld + "</p:bldLst>" if bld else ""}</p:timing>')
    _place(slide, etree.fromstring(xml), "timing")


def _has_fill(shp) -> bool:
    spPr = shp._element.find(qn("p:spPr"))
    return spPr is not None and (spPr.find(qn("a:solidFill")) is not None or spPr.find(qn("a:gradFill")) is not None)


def add_transition(slide, kind="fade", dur_ms=900):
    if kind in (None, "none"):
        return
    if kind == "morph":
        xml = (f'<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
               f'<mc:Choice xmlns:p159="http://schemas.microsoft.com/office/powerpoint/2015/09/main" Requires="p159">'
               f'<p:transition {NSDECL} xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main" '
               f'spd="slow" p14:dur="{max(dur_ms, 1200)}"><p159:morph option="byObject"/></p:transition></mc:Choice>'
               f'<mc:Fallback><p:transition {NSDECL} spd="slow"><p:fade/></p:transition></mc:Fallback>'
               f'</mc:AlternateContent>')
    else:
        inner = {"fade": "<p:fade/>", "push": '<p:push dir="u"/>', "wipe": '<p:wipe dir="r"/>'}.get(kind, "<p:fade/>")
        xml = f'<p:transition {NSDECL} spd="med">{inner}</p:transition>'
    _place(slide, etree.fromstring(xml), "transition")


def _place(slide, el, what):
    """Chèn đúng thứ tự trong p:sld: cSld, clrMapOvr, transition, timing, extLst."""
    sld = slide._element
    tags = {"transition": [qn("p:transition"), "{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent"],
            "timing": [qn("p:timing")]}
    for t in tags[what]:
        for old in sld.findall(t):
            sld.remove(old)
    ext = sld.find(qn("p:extLst"))
    if what == "transition":
        anchor = sld.find(qn("p:clrMapOvr"))
        if anchor is None:
            anchor = sld.find(qn("p:cSld"))
        anchor.addnext(el)
    else:
        if ext is not None:
            ext.addprevious(el)
        else:
            sld.append(el)
