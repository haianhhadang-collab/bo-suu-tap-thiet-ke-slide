"""Đọc .pptx / .docx / .pdf / .md / .txt thành mô hình Deck chung."""
from __future__ import annotations

import math
import re
import unicodedata
from pathlib import Path

from .model import BULLET_RE, Deck, Img, Item, Slide, clean, strip_bullet

STEPTAG_RE = re.compile(r"^\s*(?:(?:step|bước|phần|part|giai đoạn|phase|stage|module|unit|bài|chương)\s*)?\d{1,2}[.:)]?\s*$",
                        re.I)
CITE_RE = re.compile(r"(\(\s*(?:19|20)\d{2}[a-z]?\s*\)|et al\.|adapted from|nguồn\s*:|source\s*:|theo\s+[A-ZĐ])", re.I)


def read(path: str | Path) -> Deck:
    p = Path(path)
    ext = p.suffix.lower()
    if ext == ".pptx":
        return read_pptx(p)
    if ext == ".docx":
        return read_docx(p)
    if ext == ".pdf":
        return read_pdf(p)
    if ext in (".md", ".markdown"):
        return read_md(p.read_text(encoding="utf-8-sig"), p.parent)
    if ext == ".txt":
        return read_txt(p.read_text(encoding="utf-8-sig"))
    raise ValueError(f"Chưa hỗ trợ loại file {ext} (dùng .pptx, .docx, .pdf, .md, .txt)")


# ======================================================================
# PowerPoint
# ======================================================================

def _iter_shapes(shapes, dx=0, dy=0):
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(sh.shapes)
        else:
            yield sh


def read_pptx(path: Path) -> Deck:
    from pptx import Presentation
    from pptx.enum.shapes import PP_PLACEHOLDER
    prs = Presentation(str(path))
    SW, SH = prs.slide_width or 12192000, prs.slide_height or 6858000
    deck = Deck()
    raw: list[Slide] = []
    for si, sl in enumerate(prs.slides):
        frames, imgs, table = [], [], None
        title = ""
        for sh in _iter_shapes(sl.shapes):
            ph = sh.placeholder_format.type if sh.is_placeholder else None
            if ph in (PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER):
                continue
            top = (sh.top or 0) / SH
            left = (sh.left or 0) / SW
            wid = (sh.width or 0) / SW
            hei = (sh.height or 0) / SH
            blob = None
            try:
                if sh.shape_type == 13 or sh.is_placeholder:
                    blob = sh.image
            except (AttributeError, ValueError):
                blob = None
            if blob is not None:
                if wid * 13.33 < 0.6 and hei * 7.5 < 0.6:
                    continue
                im = Img.from_bytes(blob.blob, blob.ext)
                if im:
                    imgs.append((top, left, im))
                continue
            if getattr(sh, "has_table", False) and sh.has_table:
                table = [[clean(c.text) for c in r.cells] for r in sh.table.rows]
                continue
            if not getattr(sh, "has_text_frame", False) or not sh.has_text_frame:
                continue
            paras = []
            maxsz = 0
            bold = False
            for p in sh.text_frame.paragraphs:
                t = clean(p.text)
                if not t:
                    continue
                for r in p.runs:
                    if r.font.size:
                        maxsz = max(maxsz, r.font.size.pt)
                    bold = bold or bool(r.font.bold)
                paras.append((t, p.level))
            if not paras:
                continue
            is_title = ph in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
            is_sub = ph == PP_PLACEHOLDER.SUBTITLE
            frames.append(dict(top=top, left=left, w=wid, h=hei, paras=paras, size=maxsz or 18, bold=bold,
                               title=is_title, sub=is_sub))
        notes = ""
        if sl.has_notes_slide:
            notes = sl.notes_slide.notes_text_frame.text.strip()
        s = _frames_to_slide(frames, si)
        s.images = [im for _, _, im in sorted(imgs, key=lambda t: (round(t[0], 1), t[1]))]
        s.table = table
        s.notes = notes
        raw.append(s)

    if raw:
        first = raw[0]
        deck.title = first.title
        deck.subtitle = first.subtitle
        meta = [i.text for i in first.items]
        if not deck.subtitle and meta and len(meta[0]) < 120:
            deck.subtitle = meta.pop(0)
        deck.meta = meta
        deck.cover_notes = first.notes
        deck.slides = raw[1:]
    return deck


def _frames_to_slide(frames: list[dict], si: int) -> Slide:
    s = Slide()
    if not frames:
        return s
    # tiêu đề: placeholder tiêu đề, nếu không có thì chữ cỡ lớn nhất ở nửa trên
    tf = next((f for f in frames if f["title"]), None)
    if tf is None:
        cands = [f for f in frames if f["top"] < 0.42 and len(" ".join(t for t, _ in f["paras"])) <= 140]
        if cands:
            tf = max(cands, key=lambda f: (f["size"], -f["top"]))
            if tf["size"] < 20 and si > 0 and max(f["size"] for f in frames) > tf["size"]:
                tf = None
    if tf:
        s.title = " ".join(t for t, _ in tf["paras"])
        frames = [f for f in frames if f is not tf]
    sub = next((f for f in frames if f["sub"]), None)
    if sub:
        s.subtitle = " ".join(t for t, _ in sub["paras"])
        frames = [f for f in frames if f is not sub]

    # nguồn trích dẫn: chữ nhỏ ở đáy slide có dạng trích dẫn
    body_sizes = sorted(f["size"] for f in frames) or [18]
    med = body_sizes[len(body_sizes) // 2]
    keep = []
    for f in frames:
        txt = " ".join(t for t, _ in f["paras"])
        if f["top"] > 0.82 and (CITE_RE.search(txt) or f["size"] < med - 2) and len(txt) < 260:
            s.source = (s.source + " " + txt).strip()
        else:
            keep.append(f)
    frames = sorted(keep, key=lambda f: (round(f["top"] / 0.035), f["left"]))

    # phụ đề: dòng ngắn sát ngay dưới tiêu đề, cùng lề trái, đứng một mình trên hàng
    if tf and not s.subtitle and frames:
        f0 = min(frames, key=lambda f: f["top"])
        txt = " ".join(t for t, _ in f0["paras"])
        gap = f0["top"] - (tf["top"] + tf["h"])
        alone = not any(f is not f0 and abs(f["top"] - f0["top"]) < 0.03 for f in frames)
        if (len(f0["paras"]) == 1 and len(txt) < 90 and not txt.isupper() and -0.02 < gap < 0.07
                and abs(f0["left"] - tf["left"]) < 0.03 and len(frames) > 1 and alone):
            s.subtitle = txt
            frames = [f for f in frames if f is not f0]

    def text_of(f):
        return " ".join(t for t, _ in f["paras"])

    def is_short(f):
        t = text_of(f)
        return (len(f["paras"]) == 1 and not t[:1].islower()
                and (len(t.split()) <= 4 or (t.isupper() and len(t) <= 32)))

    def overlap(f, g):
        lo, hi = max(f["left"], g["left"]), min(f["left"] + f["w"], g["left"] + g["w"])
        return max(0.0, hi - lo) / max(min(f["w"], g["w"]), 1e-6)

    # 1) nhãn ngắn đứng sát nhau trên cùng một hàng => gộp (vd "STEP 1" + "NOTICE")
    merged = []
    skip = set()
    for i, f in enumerate(frames):
        if i in skip:
            continue
        if is_short(f) and STEPTAG_RE.match(text_of(f)):
            for k in range(i + 1, len(frames)):
                g = frames[k]
                if k in skip or not is_short(g):
                    continue
                if abs(g["top"] - f["top"]) < 0.015 and 0 <= g["left"] - (f["left"] + f["w"]) < 0.04:
                    f = dict(f, paras=[(text_of(f) + " · " + text_of(g), 0)],
                             w=g["left"] + g["w"] - f["left"])
                    skip.add(k)
        merged.append(f)
    frames = merged

    # 2) mỗi khối nội dung tìm nhãn ngắn gần nhất ngay phía trên (hoặc bên trái cùng hàng)
    label_of: dict[int, int] = {}
    taken = set()
    for bi, bf in enumerate(frames):
        if is_short(bf):
            continue
        best, best_gap = None, 9.0
        for li, lf in enumerate(frames):
            if li == bi or li in taken or not is_short(lf) or len(text_of(bf)) <= len(text_of(lf)):
                continue
            gap = bf["top"] - (lf["top"] + lf["h"])
            if -0.01 <= gap <= 0.22 and overlap(lf, bf) >= 0.5:
                between = any(k not in (li, bi) and overlap(frames[k], bf) > 0.3
                              and lf["top"] + lf["h"] - 0.005 < frames[k]["top"] < bf["top"] - 0.005
                              for k in range(len(frames)))
                if not between and gap < best_gap:
                    best, best_gap = li, gap
            elif (abs(bf["top"] - lf["top"]) < 0.03 and 0 <= bf["left"] - (lf["left"] + lf["w"]) < 0.1
                  and 0.5 < best_gap):
                best, best_gap = li, 0.5
        if best is not None:
            label_of[best] = bi
            taken.add(best)

    items: list[Item] = []
    consumed = set(label_of.values())
    for i, f in enumerate(frames):
        if i in consumed:
            continue
        paras = f["paras"]
        txt = text_of(f)
        if i in label_of:
            body = frames[label_of[i]]
            lab = txt.title() if txt.isupper() else txt
            if len(body["paras"]) >= 4 or len(text_of(body)) > 320:
                items.append(Item(text=lab, kind="head"))
                for t, lvl in _merge_broken(body["paras"]):
                    items.append(Item(text=strip_bullet(t), kind="bullet", level=min(lvl, 2)))
            else:
                items.append(Item(text=_join_lines(_merge_broken(body["paras"])), kind="card", label=lab))
            continue
        if is_short(f):
            items.append(Item(text=txt.title() if txt.isupper() and "·" in txt else txt, kind="chip",
                              level=round(f["top"] / 0.04)))
            continue
        if len(paras) == 1:
            t = paras[0][0]
            items.append(Item(text=strip_bullet(t), kind="bullet" if BULLET_RE.match(t) else "para"))
        else:
            mp = _merge_broken(paras)
            for t, lvl in mp:
                items.append(Item(text=strip_bullet(t), kind="bullet" if len(mp) > 1 else "para", level=min(lvl, 2)))
    _fix_chips(items)
    # đoạn ngắn kiểu tiêu đề con, theo sau là gạch đầu dòng => tiêu đề nhóm
    for k, it in enumerate(items[:-1]):
        if (it.kind == "para" and len(it.text.split()) <= 8 and not it.text.endswith((".", "!", ";", ","))
                and items[k + 1].kind == "bullet"):
            it.kind = "head"
    s.items = items
    return s


def _merge_broken(paras):
    """Nối các dòng bị ngắt giữa câu (dòng sau bắt đầu bằng chữ thường)."""
    out = []
    for t, lvl in paras:
        if out and t[:1].islower() and not BULLET_RE.match(t) and not out[-1][0].endswith((".", "!", "?", ";")):
            out[-1] = (out[-1][0] + " " + t, out[-1][1])
        else:
            out.append((t, lvl))
    return out


def _fix_chips(items: list[Item]):
    """Chip chỉ giữ khi có ít nhất 2 chip cùng hàng; chip lẻ thành tiêu đề con hoặc đoạn văn."""
    rows: dict[int, list[Item]] = {}
    for it in items:
        if it.kind == "chip" and "·" not in it.text:
            rows.setdefault(it.level, []).append(it)
    for row in rows.values():
        if len(row) < 2:
            for it in row:
                it.kind = "head" if it.text.endswith(("?", ":")) or len(it.text.split()) >= 2 else "para"
    for it in items:
        if it.kind == "chip":
            it.level = 0


def _join_lines(paras) -> str:
    ts = [strip_bullet(t) for t, _ in paras]
    return ts[0] if len(ts) == 1 else "\n".join(ts)


# ======================================================================
# Word
# ======================================================================

HEAD_STYLE_RE = re.compile(r"^(?:heading|đề mục|đầu đề|tiêu đề)\s*(\d)$")
TOC_STYLE_RE = re.compile(r"^(?:toc|mục lục)(?:\s|$)")


def read_docx(path: Path) -> Deck:
    """Tách trang theo Heading: mỗi heading (mọi cấp, kể cả kiểu tự đặt có cấp đề cương) mở một trang.
    Heading cấp ngoài cùng có mục con, chỉ có tên hoặc vài dòng dẫn: trang chương. Heading cấp trong chỉ có tên:
    không có trang riêng, tên hiện làm phụ đề các trang con.
    Mục con cấp 3 trở xuống mà mục nào cũng ngắn thì gộp thành thẻ trên trang mục cha."""
    import docx
    from docx.oxml.ns import qn
    doc = docx.Document(str(path))
    st, pdef, dsz = _docx_styles(doc.styles.element)
    deck = Deck()
    blocks: list[dict] = []
    rels = doc.part.related_parts
    for el in doc.element.body.iterchildren():
        tag = el.tag.split("}")[1]
        if tag == "tbl":
            tbl = docx.table.Table(el, doc)
            rows = [[clean(c.text) for c in r.cells] for r in tbl.rows]
            rows = [r for r in rows if any(r)]
            if rows:
                blocks.append(dict(t="table", rows=rows))
            continue
        if tag != "p":
            continue
        para = docx.text.paragraph.Paragraph(el, doc)
        ppr = el.find(qn("w:pPr"))
        ps = ppr.find(qn("w:pStyle")) if ppr is not None else None
        sid = ps.get(qn("w:val")) if ps is not None else pdef
        style = st[sid]["name"] if sid in st else (sid or "").lower()
        txt = clean(para.text)
        # ảnh trong đoạn
        for blip in el.iter(qn("a:blip")):
            rid = blip.get(qn("r:embed"))
            part = rels.get(rid)
            if part is not None:
                im = Img.from_bytes(part.blob)
                if im:
                    blocks.append(dict(t="img", img=im))
        if not txt or TOC_STYLE_RE.match(style):
            continue
        if style == "title":
            deck.title = txt
            continue
        if style == "subtitle":
            deck.subtitle = txt
            continue
        hl = _head_level(st, sid, ppr)
        if hl:
            blocks.append(dict(t="h", level=hl, text=txt))
            continue
        is_list = "list" in style or ppr is not None and ppr.find(qn("w:numPr")) is not None
        lvl = 0
        if is_list:
            ilvl = el.find(f".//{qn('w:ilvl')}")
            lvl = int(ilvl.get(qn("w:val"))) if ilvl is not None else 0
        if "quote" in style or "trích" in style:
            item = Item(text=txt, kind="quote")
        elif is_list or BULLET_RE.match(txt):
            item = Item(text=strip_bullet(txt), kind="bullet", level=min(lvl, 2))
        else:
            item = Item(text=txt, kind="para")
        bold, size = _para_fmt(el, st, sid, dsz)
        blocks.append(dict(t="p", item=item, text=txt, bold=bold, size=size, list=is_list))

    if not any(b["t"] == "h" for b in blocks):
        _guess_heads(blocks)
    _blocks_deck(deck, blocks)
    if not deck.title:
        deck.title = path.stem
    return deck


def _blocks_deck(deck: Deck, blocks: list[dict], card_level: int = 3):
    """Khối nội dung (h / p / table / img) -> cây mục theo heading -> trang bìa và các trang (dùng cho Word và PDF).
    card_level: các mục con từ cấp này trở xuống mà đều ngắn thì gộp thành thẻ."""
    root = _tree(blocks)
    pre, tops = root.body, root.kids
    if not deck.title and not any(b["t"] == "p" for b in pre) and tops:
        first = tops[0]
        if len(tops) == 1 and first.kids:   # một heading bao cả bài => tên bài
            deck.title, pre, tops = first.title, first.body, first.kids
        elif not first.kids and len(tops) > 1 and _cover_like(first.body):
            deck.title, pre, tops = first.title, first.body, tops[1:]
    slides: list[Slide] = []
    # phần trước heading đầu tiên: tên bài, phụ đề, các dòng thông tin trên trang bìa
    rest = []
    for b in pre:
        if b["t"] == "p" and not deck.title:
            deck.title = b["text"]
        elif b["t"] == "p" and not deck.subtitle and not rest and len(b["text"]) < 140:
            deck.subtitle = b["text"]
        else:
            rest.append(b)
    if rest and len(rest) <= 4 and all(b["t"] == "p" and len(b["text"]) < 140 for b in rest):
        deck.meta = [b["item"].text for b in rest]
    elif rest:
        slides += _body_slides(deck.title, "", rest)
    for n in tops:
        _emit(n, 1, "", slides, card_level)
    deck.slides = [s for s in slides if s.title or s.items or s.images or s.table]
    if not any(s.title for s in deck.slides):
        _auto_titles(deck)


def _int(el, tag):
    from docx.oxml.ns import qn
    e = el.find(qn(tag)) if el is not None else None
    v = e.get(qn("w:val")) if e is not None else None
    return int(v) if v is not None and re.fullmatch(r"-?\d+", v) else None


def _onoff(el, tag):
    from docx.oxml.ns import qn
    e = el.find(qn(tag)) if el is not None else None
    return None if e is None else e.get(qn("w:val")) not in ("0", "false", "off")


def _docx_styles(root):
    """styleId -> tên, kiểu gốc, cấp đề cương, đậm, cỡ chữ; kèm kiểu đoạn mặc định và cỡ chữ mặc định."""
    from docx.oxml.ns import qn
    st, pdef = {}, ""
    for s in root.findall(qn("w:style")):
        sid = s.get(qn("w:styleId")) or ""
        nm, based = s.find(qn("w:name")), s.find(qn("w:basedOn"))
        ppr, rpr = s.find(qn("w:pPr")), s.find(qn("w:rPr"))
        st[sid] = dict(name=((nm.get(qn("w:val")) if nm is not None else None) or sid).lower(),
                       based=(based.get(qn("w:val")) if based is not None else None) or "",
                       outline=_int(ppr, "w:outlineLvl"), bold=_onoff(rpr, "w:b"), size=_int(rpr, "w:sz"))
        if s.get(qn("w:type")) == "paragraph" and s.get(qn("w:default")) in ("1", "true"):
            pdef = sid
    dr = root.find(qn("w:docDefaults"))
    for t in ("w:rPrDefault", "w:rPr"):
        dr = dr.find(qn(t)) if dr is not None else None
    return st, pdef, _int(dr, "w:sz") or 22


def _chain(st, sid):
    """Kiểu và các kiểu gốc của nó (basedOn)."""
    seen = []
    while sid and sid in st and sid not in seen and len(seen) < 12:
        seen.append(sid)
        sid = st[sid]["based"]
    return [st[s] for s in seen]


def _head_level(st, sid, ppr) -> int:
    """Cấp heading 1–9 (0 = đoạn thường): cấp đề cương của đoạn, rồi tên kiểu / cấp đề cương của kiểu và kiểu gốc."""
    o = _int(ppr, "w:outlineLvl")
    if o is not None:
        return o + 1 if o < 9 else 0
    for s in _chain(st, sid):
        m = HEAD_STYLE_RE.match(s["name"])
        if m:
            return int(m.group(1))
        if s["outline"] is not None:
            return s["outline"] + 1 if s["outline"] < 9 else 0
    return 0


def _para_fmt(el, st, sid, dsz):
    """(cả đoạn in đậm, cỡ chữ lớn nhất theo nửa point)."""
    from docx.oxml.ns import qn
    pc = _chain(st, sid)
    pb = next((s["bold"] for s in pc if s["bold"] is not None), False)
    ps = next((s["size"] for s in pc if s["size"]), None) or dsz
    allb, anyt, size = True, False, 0
    for r in el.iter(qn("w:r")):
        if not "".join(x.text or "" for x in r.iter(qn("w:t"))).strip():
            continue
        anyt = True
        rpr = r.find(qn("w:rPr"))
        b, sz = _onoff(rpr, "w:b"), _int(rpr, "w:sz")
        rs = rpr.find(qn("w:rStyle")) if rpr is not None else None
        if rs is not None:
            cc = _chain(st, rs.get(qn("w:val")))
            if b is None:
                b = next((s["bold"] for s in cc if s["bold"] is not None), None)
            if not sz:
                sz = next((s["size"] for s in cc if s["size"]), None)
        if b is None:
            b = pb
        allb = allb and b
        size = max(size, sz or ps)
    return anyt and allb, size


def _num_depth(t: str) -> int:
    """Độ sâu theo cách đánh số đầu dòng: Chương/Phần, I. => 0; 1. => 1; 1.1. => 2; a) => 5."""
    if re.match(r"(?:chương|phần|part|chapter)\s+[\dIVXLC]+", t, re.I) or re.match(r"[IVXLC]+[.)]\s", t):
        return 0
    m = re.match(r"(\d+(?:\.\d+)*)\.?\)?\s", t)
    if m:
        return m.group(1).count(".") + 1
    if re.match(r"[a-zđ][.)]\s", t):
        return 5
    return 0 if t.isupper() else 1


def _guess_heads(blocks):
    """Văn bản không dùng kiểu Heading: dòng in đậm cả dòng, ngắn, không kết thúc bằng dấu chấm, phẩy => tiêu đề.
    Cấp xếp theo cỡ chữ (to trước) rồi theo cách đánh số."""
    paras = [b for b in blocks if b["t"] == "p"]
    cands = [b for b in paras if b["bold"] and not b["list"] and len(b["text"]) <= 100
             and len(b["text"].split()) <= 14 and not re.search(r"[.;,]$", b["text"])]
    if len(cands) < 2 or len(cands) > 0.5 * len(paras):
        return
    keys = sorted({(-b["size"], _num_depth(b["text"])) for b in cands})
    for b in cands:
        b["level"] = keys.index((-b["size"], _num_depth(b["text"]))) + 1
        b["t"] = "h"
        b["text"] = re.sub(r"\s*[:：]$", "", b["text"])


class _Sec:
    """Một mục của văn bản: tên (heading), nội dung ngay dưới nó, các mục con."""

    def __init__(self, level: int, title: str):
        self.level, self.title = level, title
        self.body: list[dict] = []
        self.kids: list[_Sec] = []


def _tree(blocks) -> _Sec:
    root = _Sec(0, "")
    stack = [root]
    for b in blocks:
        if b["t"] == "h":
            while stack[-1].level >= b["level"]:
                stack.pop()
            n = _Sec(b["level"], b["text"])
            stack[-1].kids.append(n)
            stack.append(n)
        else:
            stack[-1].body.append(b)
    return root


def _cover_like(body) -> bool:
    return len(body) <= 3 and all(b["t"] == "p" and len(b["text"]) < 140 for b in body)


def _intro(body) -> bool:
    """Lời dẫn ngắn của chương (tối đa 2 đoạn, 220 ký tự): hiện dưới tên chương trên trang chương."""
    return len(body) <= 2 and all(b["t"] == "p" for b in body) and sum(len(b["item"].text) for b in body) <= 220


def _body_slides(title, sub, blocks) -> list[Slide]:
    cur = Slide(title=title, subtitle=sub)
    out = [cur]
    for b in blocks:
        if b["t"] == "table":
            if cur.items or cur.table or cur.images:
                cur = Slide(title=title, subtitle=sub, cont=bool(title))
                out.append(cur)
            cur.table = b["rows"]
        elif b["t"] == "img":
            cur.images.append(b["img"])
        else:
            cur.items.append(b["item"])
    return out


def _short(n: _Sec) -> bool:
    """Mục ngắn: không có mục con, bảng, ảnh; tối đa 4 ý, tổng 300 ký tự."""
    if n.kids or any(b["t"] != "p" for b in n.body):
        return False
    ts = [b["item"].text for b in n.body]
    return len(ts) <= 4 and sum(len(t) for t in ts) <= 300 and (len(ts) <= 1 or max(len(t) for t in ts) <= 120)


def _emit(n: _Sec, depth: int, crumb: str, out: list[Slide], card_level: int = 3):
    kids = n.kids
    if len(kids) >= 2 and (depth >= 2 or min(k.level for k in kids) >= card_level) and all(_short(k) for k in kids):
        ss = _body_slides(n.title, crumb, n.body)
        if ss[-1].table:
            ss.append(Slide(title=n.title, subtitle=crumb, cont=True))
        if all(not k.body for k in kids):
            ss[-1].items += [Item(text=k.title, kind="bullet") for k in kids]
        else:
            ss[-1].items += [Item(text="\n".join(b["item"].text for b in k.body), kind="card", label=k.title)
                             for k in kids]
        out += ss
        return
    sub = ""
    if depth == 1 and kids and _intro(n.body):   # chương chỉ có tên hoặc vài dòng dẫn => trang chương
        out.append(Slide(title=n.title, hint="section", items=[b["item"] for b in n.body]))
    elif n.body or not kids:
        out += _body_slides(n.title, crumb, n.body)
    else:   # mục trung gian chỉ có tên: hiện làm phụ đề trên các trang con
        sub = f"{crumb} · {n.title}" if crumb else n.title
    for k in kids:
        _emit(k, depth + 1, sub, out, card_level)


# ======================================================================
# PDF có lớp chữ
# ======================================================================

SPACE_GAP = 0.102   # khoảng trống giữa hai chữ (theo cỡ chữ) coi là dấu cách, như pdf.js
BOLD_FONT_RE = re.compile(r"bold|black|heavy|semibold|demibold", re.I)
ITALIC_FONT_RE = re.compile(r"italic|oblique", re.I)
PDF_BULLET_RE = re.compile(r"^[-•●▪■◦○◆►▶✓✔→]\s*")
PAGENUM_RE = re.compile(r"^(?:trang|page)?\s*#(?:\s*(?:/|của|of)\s*#)?$|^-\s*#\s*-$")
TOC_TAIL_RE = re.compile(r"^(.*?)[\s.…·_]+\d{1,4}$")
TOC_TITLE_RE = re.compile(r"^(?:mục lục|contents|table of contents)$", re.I)
PDF_PASS_MSG = "File PDF có mật khẩu. Hãy lưu lại một bản không đặt mật khẩu rồi thử lại."
PDF_SCAN_MSG = ("File PDF này là ảnh chụp hoặc bản scan, không có lớp chữ để đọc. "
                "Công cụ chưa hỗ trợ loại này (cần nhận dạng chữ OCR).")
PDF_EMPTY_MSG = "Không tìm thấy chữ nào trong file PDF."


def _half(v: float) -> float:
    return math.floor(v * 2 + 0.5) / 2


def _norm(t: str) -> str:
    return re.sub(r"\s+", " ", unicodedata.normalize("NFC", t)).strip().lower()


def read_pdf(path: Path) -> Deck:
    """PDF có lớp chữ: dựng lại dòng, đoạn, gạch đầu dòng, bảng, ảnh; heading lấy từ mục lục điều hướng (bookmark),
    không có thì đoán theo cỡ chữ, chữ đậm; rồi tách trang theo heading như file Word.
    PDF xuất từ PowerPoint (trang ngang, không bookmark): mỗi trang một slide, dòng chữ to nhất là tiêu đề."""
    import pymupdf
    doc = pymupdf.open(str(path))
    if doc.needs_pass:
        raise ValueError(PDF_PASS_MSG)
    pages = []
    for pg in doc:
        pieces = []
        for b in pg.get_text("rawdict")["blocks"]:
            for ln in b.get("lines", []):
                if abs(ln["dir"][1]) > 0.01 or ln["dir"][0] <= 0:   # bỏ chữ xoay
                    continue
                pend = False   # có dấu cách đứng trước mẩu kế tiếp
                for sp in ln["spans"]:
                    size, y = _half(sp["size"]), sp["origin"][1]
                    bold, it = bool(BOLD_FONT_RE.search(sp["font"])), bool(ITALIC_FONT_RE.search(sp["font"]))
                    cur, last = None, 0.0
                    for ch in sp["chars"]:
                        x0, x1 = ch["bbox"][0], ch["bbox"][2]
                        if cur and x0 - last > SPACE_GAP * size:   # khoảng trống giữa hai ký tự: tách mẩu như pdf.js
                            pieces.append(_piece(cur))
                            cur = None
                        last = x1
                        if ch["c"].isspace():
                            if cur:
                                cur["text"] += " "
                            else:
                                pend = True
                            continue
                        if cur is None:
                            cur = dict(text=ch["c"], x0=x0, x1=x1, y=y, size=size, bold=bold, italic=it, pre=pend)
                            pend = False
                        else:
                            cur["text"] += ch["c"]
                            cur["x1"] = x1
                    if cur:
                        pieces.append(_piece(cur))
        W, H = pg.rect.width, pg.rect.height
        imgs, seen = [], set()
        for info in pg.get_image_info():
            x0, y0, x1, y1 = info["bbox"]
            x0, y0, x1, y1 = max(x0, 0), max(y0, 0), min(x1, W), min(y1, H)
            key = (round(x0), round(y0), round(x1), round(y1))
            # bỏ ảnh nhỏ, ảnh nền cả trang, dải trang trí dài mảnh
            if (x1 - x0 >= 40 and y1 - y0 >= 40 and (x1 - x0) * (y1 - y0) <= 0.85 * W * H
                    and max(x1 - x0, y1 - y0) <= 6 * min(x1 - x0, y1 - y0) and key not in seen):
                seen.add(key)
                imgs.append(dict(x0=x0, y0=y0, x1=x1, y1=y1, get=lambda pg=pg, r=(x0, y0, x1, y1): Img.from_bytes(
                    pg.get_pixmap(clip=pymupdf.Rect(*r), dpi=150).tobytes("png"), "png")))
        pages.append(dict(w=W, h=H, pieces=pieces, imgs=imgs))
    outline = [(lvl, title, p - 1) for lvl, title, p in doc.get_toc(simple=True)]
    blocks = _pdf_blocks(pages, outline)
    for b in blocks:
        if b["t"] == "img":
            b["img"] = b.pop("get")()
    deck = Deck()
    _blocks_deck(deck, [b for b in blocks if b["t"] != "img" or b["img"]], card_level=2)
    if not deck.title:
        deck.title = path.stem
    return deck


def _piece(p):
    """Mẩu chữ: bỏ dấu cách hai đầu, nhớ lại là có dấu cách đứng trước / sau."""
    p["post"] = p["text"].endswith(" ")
    p["text"] = p["text"].strip()
    return p


def _pdf_lines(pieces):
    """Mẩu chữ -> dòng (cùng đường chân chữ) -> các đoạn trong dòng (cách nhau quá 2 lần cỡ chữ: cột, ô bảng)."""
    rows = []
    for p in sorted(pieces, key=lambda p: (p["y"], p["x0"])):
        r = rows[-1] if rows else None
        if r and abs(p["y"] - r["y"]) <= 0.3 * max(p["size"], r["size"]):
            r["ps"].append(p)
            r["size"] = max(r["size"], p["size"])
        else:
            rows.append(dict(y=p["y"], size=p["size"], ps=[p]))
    out = []
    for r in rows:
        segs, cur = [], None
        for p in sorted(r["ps"], key=lambda p: p["x0"]):
            gap = p["x0"] - cur["x1"] if cur else 0
            if cur and gap > 2 * max(p["size"], cur["size"]):
                segs.append(cur)
                cur = None
            if cur is None:
                cur = dict(p)
                continue
            if cur["post"] or p["pre"] or gap > SPACE_GAP * p["size"]:
                cur["text"] += " "
            cur["text"] += p["text"]
            cur["post"] = p["post"]
            cur["x1"] = max(cur["x1"], p["x1"])
            cur["size"] = max(cur["size"], p["size"])
            cur["bold"] = cur["bold"] and p["bold"]
            cur["italic"] = cur["italic"] and p["italic"]
        segs.append(cur)
        for s in segs:
            s["text"] = clean(unicodedata.normalize("NFC", s["text"]))
        segs = [s for s in segs if s["text"]]
        if segs:
            out.append(dict(y=r["y"], size=max(s["size"] for s in segs), bold=all(s["bold"] for s in segs),
                            italic=all(s["italic"] for s in segs), segs=segs,
                            x0=segs[0]["x0"], x1=max(s["x1"] for s in segs), text=" ".join(s["text"] for s in segs)))
    return out


def _pdf_blocks(pages, outline) -> list[dict]:
    """Các trang PDF (mẩu chữ, khung ảnh) -> khối nội dung h / p / table / img như file Word."""
    n = len(pages)
    for i, pg in enumerate(pages):
        pg["lines"] = _pdf_lines(pg["pieces"])
        for l in pg["lines"]:
            l["page"] = i
    chars = sum(len(l["text"]) for pg in pages for l in pg["lines"])
    if not chars or chars < 10 * n and any(pg["imgs"] for pg in pages):
        raise ValueError(PDF_SCAN_MSG if any(pg["imgs"] for pg in pages) else PDF_EMPTY_MSG)

    # đầu trang, chân trang: lặp lại ở nửa số trang trở lên, hoặc chỉ là số trang
    def hf(t):
        return re.sub(r"\d+", "#", t.lower()).strip()

    def edge(l, pg):
        return l["y"] < 0.08 * pg["h"] or l["y"] > 0.93 * pg["h"]

    cnt: dict[str, int] = {}
    for pg in pages:
        for k in {hf(l["text"]) for l in pg["lines"] if edge(l, pg)}:
            cnt[k] = cnt.get(k, 0) + 1
    for pg in pages:
        pg["lines"] = [l for l in pg["lines"] if not (edge(l, pg) and (
            PAGENUM_RE.match(hf(l["text"])) or n >= 2 and cnt[hf(l["text"])] >= max(2, 0.5 * n)))]
    L = [l for pg in pages for l in pg["lines"]]
    for l in L:
        if TOC_TITLE_RE.match(l["text"]):
            l["drop"] = True
    sizes: dict[float, int] = {}
    for l in L:
        sizes[l["size"]] = sizes.get(l["size"], 0) + len(l["text"])
    body = max(sizes, key=lambda s: (sizes[s], -s))

    # bảng: từ 2 dòng liền nhau trở lên có cùng số cột, các cột thẳng hàng
    tid = 0
    for pg in pages:
        ls, i = pg["lines"], 0
        while i < len(ls):
            j = i + 1
            if len(ls[i]["segs"]) >= 2:
                while (j < len(ls) and len(ls[j]["segs"]) == len(ls[i]["segs"])
                       and ls[j]["y"] - ls[j - 1]["y"] <= 3 * max(ls[j]["size"], ls[j - 1]["size"])
                       and all(abs(s["x0"] - t["x0"]) <= max(4, 0.8 * ls[i]["size"])
                               for s, t in zip(ls[i]["segs"], ls[j]["segs"]))):
                    j += 1
                if j - i >= 2:
                    tid += 1
                    for k in range(i, j):
                        ls[k]["table"] = tid
            i = j

    # heading: theo mục lục điều hướng (bookmark) của file
    found = 0
    i0 = 0
    for lvl, title, pno in outline:
        nt = _norm(title)
        hit = _pdf_find(L, i0, nt, pno) if nt else None
        if hit:
            i, k = hit
            L[i]["head"] = lvl
            L[i]["text"] = " ".join(L[j]["text"] for j in range(i, k + 1))
            for j in range(i + 1, k + 1):
                L[j]["drop"] = True
            i0, found = k + 1, found + 1
    if not found and all(pg["w"] > 1.2 * pg["h"] for pg in pages):
        # trang chiếu (PDF xuất từ PowerPoint): dòng chữ to trên cùng mỗi trang là tiêu đề (bỏ số chương "01")
        for p, pg in enumerate(pages):
            ls = [l for l in pg["lines"] if not l.get("table") and not l.get("drop")]
            s = next((k for k, l in enumerate(ls) if l["size"] >= body + 2 and not re.fullmatch(r"\d{1,3}", l["text"])), None)
            if s is None:
                continue
            top = ls[s]["size"]
            e = s + 1
            while e < len(ls) and ls[e]["size"] == top and ls[e]["y"] - ls[e - 1]["y"] <= 1.6 * top:
                e += 1
            if p:
                ls[s]["head"] = 1
            else:   # trang đầu là trang bìa: tiêu đề thành tên bài, các dòng sau thành phụ đề, thông tin
                ls[s]["cover"] = True
            ls[s]["key"] = -1.0   # tiêu đề đứng đầu trang, trước cả các dòng nằm trên nó
            ls[s]["text"] = " ".join(l["text"] for l in ls[s:e])
            for l in ls[s + 1:e]:
                l["drop"] = True
            for l in ls[:s]:
                if re.fullmatch(r"\d{1,3}", l["text"]):   # số chương "01" trang trí
                    l["drop"] = True
    elif not found:
        _pdf_guess_heads(L, body)
    heads = {_norm(l["text"]) for l in L if l.get("head")} | {_norm(t) for _, t, _ in outline}
    for l in L:   # dòng mục lục: tên heading + số trang
        m = TOC_TAIL_RE.match(l["text"])
        if not l.get("head") and m and _norm(m.group(1)) in heads:
            l["drop"] = True

    # dựng đoạn, gạch đầu dòng, bảng, ảnh theo thứ tự trên trang
    blocks: list[dict] = []
    cur = None

    def flush():
        nonlocal cur
        if cur:
            blocks.append(dict(t="p", item=Item(text=cur["text"], kind=cur["kind"], level=cur["level"]), text=cur["text"],
                               bold=cur["bold"], size=cur["size"], list=cur["kind"] == "bullet"))
        cur = None

    for pg in pages:
        pg["imgs"] = [im for im in pg["imgs"] if sum(
            1 for l in pg["lines"] for g in l["segs"]
            if im["x0"] - 2 <= g["x0"] and g["x1"] <= im["x1"] + 2 and im["y0"] <= l["y"] <= im["y1"]) < 2]
        seq = [(l.get("key", l["y"]), 0, l) for l in pg["lines"] if not l.get("drop")] + [(im["y1"], 1, im) for im in pg["imgs"]]
        seq.sort(key=lambda t: (t[0], t[1]))
        rmax = max((l["x1"] for l in pg["lines"]), default=pg["w"])
        lmin = min((l["x0"] for l in pg["lines"]), default=0)
        for _, kind, e in seq:
            if kind:
                flush()
                blocks.append(dict(t="img", get=e["get"]))
                continue
            l = e
            if l.get("head"):
                flush()
                blocks.append(dict(t="h", level=l["head"], text=l["text"]))
                continue
            if l.get("table"):
                flush()
                row = [s["text"] for s in l["segs"]]
                tb = blocks[-1] if blocks and blocks[-1]["t"] == "table" and blocks[-1]["tid"] == l["table"] else None
                if tb and l["y"] - tb["y"] <= 1.35 * l["size"] and all(c[:1].islower() for c in row):
                    tb["rows"][-1] = [a + " " + c for a, c in zip(tb["rows"][-1], row)]   # chữ trong ô xuống dòng
                elif tb:
                    tb["rows"].append(row)
                else:
                    tb = dict(t="table", tid=l["table"], rows=[row])
                    blocks.append(tb)
                tb["y"] = l["y"]
                continue
            bt = _pdf_bullet(l["text"])
            if cur and bt is None and not l.get("cover") and not cur["cover"] and _pdf_cont(cur, l, rmax):
                cur["text"] += " " + l["text"]
                cur["last"] = l
                cur["bold"] = cur["bold"] and l["bold"]
                continue
            flush()
            d = l["x0"] - lmin
            cur = dict(kind="para" if bt is None else "bullet", text=l["text"] if bt is None else bt,
                       level=0 if bt is None or d < 2.5 * l["size"] else (1 if d < 5.5 * l["size"] else 2),
                       bold=l["bold"], size=l["size"], first=l, last=l, cover=bool(l.get("cover")))
    flush()
    for b in blocks:
        if b["t"] == "table" and len(b["rows"]) < 2:   # bảng còn một dòng (phần còn lại là mục lục) => đoạn thường
            t = " ".join(b["rows"][0])
            b.update(t="p", item=Item(text=t, kind="para"), text=t, bold=False, size=body, list=False)
    return blocks


def _pdf_find(L, i0, nt, pno):
    """Tìm dòng (hoặc 2–3 dòng liền nhau) có chữ đúng bằng tên mục trong bookmark, từ trang ghi trong bookmark."""
    for i in range(i0, len(L)):
        l = L[i]
        if pno >= 0 and l["page"] < pno:
            continue
        if pno >= 0 and l["page"] > pno + 1:
            return None
        if l.get("table") or l.get("drop"):
            continue
        acc = ""
        for k in range(i, min(i + 3, len(L))):
            acc = (acc + " " + _norm(L[k]["text"])).strip()
            if acc == nt:
                return i, k
            if not nt.startswith(acc):
                break
    return None


def _pdf_guess_heads(L, body):
    """Không có bookmark: dòng ngắn, chữ to hơn chữ thân bài hoặc in đậm, không kết thúc bằng dấu chấm, phẩy => heading.
    Dòng chữ to đứng đầu trang 1 với kiểu chữ chỉ xuất hiện một lần => tên bài, phụ đề (trang bìa)."""
    def cand(l):
        return (not l.get("table") and not l.get("drop") and len(l["text"]) <= 120 and len(l["text"].split()) <= 16
                and not re.search(r"[.;,]$", l["text"])
                and (l["size"] >= body + 1.5 or l["bold"] and l["size"] >= body - 0.5))

    groups, i = [], 0
    while i < len(L):
        if not cand(L[i]):
            i += 1
            continue
        j = i + 1
        while (j < len(L) and cand(L[j]) and L[j]["page"] == L[i]["page"] and L[j]["size"] == L[i]["size"]
               and L[j]["bold"] == L[i]["bold"] and L[j]["italic"] == L[i]["italic"] and L[j]["y"] - L[j - 1]["y"] <= 1.6 * L[i]["size"]
               and sum(len(L[k]["text"]) for k in range(i, j + 1)) <= 150):
            j += 1
        groups.append((i, j - 1))
        i = j
    plain = sum(1 for l in L if not l.get("table") and not l.get("drop"))
    if sum(1 for a, _ in groups if L[a]["size"] < body + 1.5) > 0.4 * plain:   # in đậm tràn lan: chỉ tin cỡ chữ
        groups = [g for g in groups if L[g[0]]["size"] >= body + 1.5]
    style: dict[tuple, int] = {}
    for a, _ in groups:
        k = (L[a]["size"], L[a]["bold"], L[a]["italic"])
        style[k] = style.get(k, 0) + 1
    first = next((n for n, (a, _) in enumerate(groups) if style[(L[a]["size"], L[a]["bold"], L[a]["italic"])] >= 2),
                 len(groups))
    heads = []
    for n, (a, z) in enumerate(groups):
        L[a]["text"] = " ".join(L[k]["text"] for k in range(a, z + 1))
        for k in range(a + 1, z + 1):
            L[k]["drop"] = True
        if n < first and L[a]["page"] == 0:
            L[a]["cover"] = True
        else:
            heads.append(L[a])
    keys = sorted({(-l["size"], _num_depth(l["text"]), l["italic"]) for l in heads})
    for l in heads:
        l["head"] = keys.index((-l["size"], _num_depth(l["text"]), l["italic"])) + 1


def _pdf_bullet(t: str) -> str | None:
    """Chữ sau dấu đầu dòng, hoặc None nếu dòng không bắt đầu bằng dấu đầu dòng."""
    if PDF_BULLET_RE.match(t):
        return PDF_BULLET_RE.sub("", t, count=1).strip()
    if BULLET_RE.match(t):
        return strip_bullet(t)
    return None


def _pdf_cont(cur, l, rmax) -> bool:
    """Dòng l có viết tiếp đoạn đang dựng không."""
    a = cur["last"]
    if abs(l["size"] - cur["size"]) > 1:
        return False
    if l["page"] != a["page"]:
        return l["page"] == a["page"] + 1 and not re.search(r"[.!?:;]$", cur["text"])
    gap = l["y"] - a["y"]
    if gap <= 0 or gap > 1.6 * max(l["size"], a["size"]):
        return False
    if re.search(r"[.!?:;]$", a["text"]) and a["x1"] < rmax - 2 * a["size"]:   # dòng trước ngắn và hết câu
        return False
    if cur["kind"] == "bullet":
        return l["x0"] >= cur["first"]["x0"] + 0.3 * l["size"]
    return l["x0"] <= a["x0"] + 0.8 * l["size"]


# ======================================================================
# Markdown / văn bản
# ======================================================================

def read_md(text: str, base: Path | None = None) -> Deck:
    deck = Deck()
    slides: list[Slide] = []
    cur: Slide | None = None
    label = ""
    para_buf: list[str] = []
    table_buf: list[list[str]] = []

    def flush():
        nonlocal para_buf, table_buf
        if para_buf:
            t = " ".join(para_buf)
            if cur is None:
                if not deck.subtitle:
                    deck.subtitle = t
                else:
                    deck.meta.append(t)
            elif label_used(t):
                pass
            else:
                cur.items.append(Item(text=t, kind="para"))
            para_buf = []
        if table_buf:
            rows = [r for r in table_buf if not all(re.fullmatch(r":?-{2,}:?", c.strip()) for c in r if c.strip())]
            if cur is None:
                new(deck.title)
            cur.table = rows
            table_buf = []

    def label_used(t):
        nonlocal label
        if label:
            cur.items.append(Item(text=t, kind="card", label=label))
            label = ""
            return True
        return False

    def new(title, hint=""):
        nonlocal cur
        cur = Slide(title=title, hint=hint)
        slides.append(cur)

    for raw in text.splitlines():
        line = raw.rstrip()
        st = line.strip()
        if not st:
            flush()
            continue
        if st.startswith("|") and st.endswith("|"):
            if para_buf:
                flush()
            table_buf.append([clean(c) for c in st.strip("|").split("|")])
            continue
        elif table_buf:
            flush()
        m = re.match(r"^(#{1,6})\s+(.*)$", st)
        if m:
            flush()
            lvl, t = len(m.group(1)), clean(m.group(2))
            if lvl == 1:
                if not deck.title and not slides:
                    deck.title = t
                else:
                    new(t, "section")
            elif lvl == 2:
                new(t)
            else:
                label = t
            continue
        if st in ("---", "***", "___"):
            flush()
            if cur is not None and (cur.items or cur.images):
                new(cur.title)
                cur.cont = True
            continue
        mi = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)", st)
        if mi:
            flush()
            src = mi.group(2).split(" ")[0]
            fp = (base / src) if base else Path(src)
            if fp.exists():
                im = Img.from_bytes(fp.read_bytes(), fp.suffix.lstrip("."), mi.group(1))
                if im:
                    if cur is None:
                        new(deck.title)
                    cur.images.append(im)
            continue
        low = st.lower()
        if low.startswith(("ghi chú:", "notes:", "note:", "???")):
            flush()
            if cur is not None:
                cur.notes = (cur.notes + "\n" + st.split(":", 1)[-1].strip()).strip()
            else:
                deck.cover_notes = st.split(":", 1)[-1].strip()
            continue
        if st.startswith(">"):
            flush()
            if cur is None:
                new(deck.title)
            cur.items.append(Item(text=clean(st.lstrip("> ")), kind="quote"))
            continue
        if BULLET_RE.match(st):
            flush()
            indent = len(line) - len(line.lstrip(" \t"))
            lvl = min(indent // 2, 2)
            t = _md_inline(strip_bullet(st))
            if cur is None:
                if deck.title:
                    deck.meta.append(t)
                    continue
                new("")
            if label:
                cur.items.append(Item(text=t, kind="card", label=label))
                label = ""
            else:
                cur.items.append(Item(text=t, kind="bullet", level=lvl))
            continue
        para_buf.append(_md_inline(st))
    flush()
    if label and cur is not None:
        cur.items.append(Item(text=label, kind="para"))
    deck.slides = [s for s in slides if s.title or s.items or s.images or s.table]
    if not deck.title:
        deck.title = deck.slides[0].title if deck.slides else "Bài trình bày"
    return deck


def _md_inline(t: str) -> str:
    t = re.sub(r"\*\*(.+?)\*\*|__(.+?)__", lambda m: m.group(1) or m.group(2), t)
    t = re.sub(r"(?<!\w)[*_](.+?)[*_](?!\w)", r"\1", t)
    t = re.sub(r"`([^`]+)`", r"\1", t)
    t = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", t)
    return clean(t)


def read_txt(text: str) -> Deck:
    """Văn bản trơn: dòng đầu là tên bài; mỗi khối cách nhau dòng trống là một trang
    (dòng đầu khối ngắn, không có dấu chấm cuối = tiêu đề trang)."""
    blocks = [b.strip("\n") for b in re.split(r"\n\s*\n", text.strip()) if b.strip()]
    deck = Deck()
    if not blocks:
        return deck
    head = [clean(l) for l in blocks[0].splitlines() if l.strip()]
    deck.title = head[0]
    if len(head) > 1:
        deck.subtitle = head[1]
        deck.meta = head[2:]
    for b in blocks[1:]:
        lines = [l.rstrip() for l in b.splitlines() if l.strip()]
        first = clean(lines[0])
        s = Slide()
        if len(lines) > 1 and len(first.split()) <= 12 and not first.endswith((".", "!", "?", ":", ";", ",")) \
                and not BULLET_RE.match(first):
            s.title = first
            lines = lines[1:]
        for l in lines:
            t = clean(l)
            if BULLET_RE.match(l.strip()):
                lvl = min((len(l) - len(l.lstrip())) // 2, 2)
                s.items.append(Item(text=strip_bullet(t), kind="bullet", level=lvl))
            else:
                s.items.append(Item(text=t, kind="para"))
        if len(lines) == 1 and not s.title and len(first.split()) <= 8:
            s.title, s.items = first, []
        deck.slides.append(s)
    if not any(s.title for s in deck.slides):
        _auto_titles(deck)
    return deck


def _auto_titles(deck: Deck):
    """Không có tiêu đề nào: lấy câu đầu (rút gọn) làm tiêu đề."""
    for s in deck.slides:
        if not s.title and s.items:
            t = s.items[0].text
            first = re.split(r"(?<=[.!?:])\s", t)[0]
            if len(first) <= 70 and len(s.items) > 1:
                s.title = first.rstrip(".:")
                s.items = s.items[1:]
