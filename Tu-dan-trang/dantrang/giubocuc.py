"""Giữ nguyên bố cục của file PowerPoint gốc, chỉ khoác mẫu thiết kế lên.

Với file .pptx, công cụ không dàn lại trang mà mở chính file gốc rồi:
  • giữ nguyên vị trí, kích thước, thứ tự lớp của mọi hộp chữ, bảng, ảnh, hình khối, nhóm hình;
  • đổi nền trang, màu chữ, màu hình, màu ô bảng, phông chữ theo mẫu (cả bảng màu và phông của theme);
  • vẽ trang trí, họa tiết của mẫu ở lớp dưới cùng; phần nào chạm vào chữ thì làm mờ đi;
  • phông mới rộng hơn làm chữ xuống thêm dòng thì cỡ chữ nhỏ lại vừa đủ để giữ số dòng như cũ;
  • thêm animation cho trang chưa có (trang đã có animation thì giữ nguyên) và hiệu ứng chuyển trang.
Bản web (web/giu-bo-cuc.js) làm đúng các bước này.
"""
from __future__ import annotations

import math
import re
from functools import lru_cache
from pathlib import Path

from lxml import etree
from PIL import ImageFont
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

from . import decor, readers
from . import draw as D
from . import textfit as TF
from .animate import add_animations, add_transition
from .model import Slide

A_NS = D.A_NS
EMU = D.EMU
CLOSING_RE = re.compile(r"(thank|cảm ơn|cám ơn|q\s*&\s*a|hỏi đáp|questions|xin chân thành)", re.I)
SYMBOL_FONT_RE = re.compile(r"^(wingdings|webdings|symbol|marlett|mt extra|segoe (ui symbol|ui emoji|mdl2)|font ?awesome|material)", re.I)
FILLS = ("noFill", "solidFill", "gradFill", "blipFill", "pattFill", "grpFill")
COLORS = ("srgbClr", "schemeClr", "sysClr", "prstClr", "scrgbClr", "hslClr")
PRST = {"black": "000000", "white": "FFFFFF", "red": "FF0000", "green": "008000", "blue": "0000FF",
        "yellow": "FFFF00", "gray": "808080", "grey": "808080", "darkGray": "A9A9A9", "lightGray": "D3D3D3",
        "orange": "FFA500", "darkRed": "8B0000", "darkBlue": "00008B", "darkGreen": "006400", "navy": "000080",
        "purple": "800080", "silver": "C0C0C0", "maroon": "800000", "cyan": "00FFFF", "magenta": "FF00FF"}
SCHEME_KEYS = ("dk1", "lt1", "dk2", "lt2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6",
               "hlink", "folHlink")


def _ln(el) -> str:
    return etree.QName(el).localname if isinstance(el.tag, str) else ""


def _r(v) -> int:
    """Làm tròn 0.5 lên như Math.round của bản web."""
    return math.floor(v + 0.5)


def _a(tag: str) -> str:
    return f"{{{A_NS}}}{tag}"


# ======================================================================
# màu: đọc, quy đổi, trộn
# ======================================================================

def _rgb(h: str):
    h = h.lstrip("#")
    return [int(h[i:i + 2], 16) for i in (0, 2, 4)]


def _hex(c) -> str:
    return "".join(f"{max(0, min(255, math.floor(v + 0.5))):02X}" for v in c)


def lightness(h: str) -> float:
    c = _rgb(h)
    return (max(c) + min(c)) / 510


def chroma(h: str) -> float:
    c = _rgb(h)
    return (max(c) - min(c)) / 255


def _to_hls(r, g, b):
    mx, mn = max(r, g, b), min(r, g, b)
    l = (mx + mn) / 2
    if mx == mn:
        return 0.0, l, 0.0
    d = mx - mn
    s = d / (mx + mn) if l <= 0.5 else d / (2 - mx - mn)
    rc, gc, bc = (mx - r) / d, (mx - g) / d, (mx - b) / d
    h = (bc - gc) if r == mx else (2 + rc - bc) if g == mx else (4 + gc - rc)
    return (h / 6) % 1.0, l, s


def _from_hls(h, l, s):
    if s == 0:
        return l, l, l
    m2 = l * (1 + s) if l <= 0.5 else l + s - l * s
    m1 = 2 * l - m2

    def v(t):
        t %= 1.0
        if t < 1 / 6:
            return m1 + (m2 - m1) * t * 6
        if t < 0.5:
            return m2
        if t < 2 / 3:
            return m1 + (m2 - m1) * (2 / 3 - t) * 6
        return m1
    return v(h + 1 / 3), v(h), v(h - 1 / 3)


def hue_bucket(h: str) -> int:
    """12 nhóm sắc độ (mỗi nhóm 30°), đỏ ở nhóm 0."""
    r, g, b = (v / 255 for v in _rgb(h))
    return int(((_to_hls(r, g, b)[0] * 360 + 15) % 360) // 30)


def _mods(base: str, el):
    """Áp các biến đổi màu của OOXML (lumMod, lumOff, tint, shade, satMod, alpha)."""
    r, g, b = (v / 255 for v in _rgb(base))
    alpha = None
    for m in el:
        tag = _ln(m)
        try:
            v = int(m.get("val", "0")) / 100000
        except ValueError:
            continue
        if tag == "alpha":
            alpha = v
        elif tag in ("lumMod", "lumOff", "satMod"):
            hh, l, s = _to_hls(r, g, b)
            if tag == "lumMod":
                l *= v
            elif tag == "lumOff":
                l += v
            else:
                s *= v
            r, g, b = _from_hls(hh, max(0.0, min(1.0, l)), max(0.0, min(1.0, s)))
        elif tag == "tint":
            r, g, b = (c * v + (1 - v) for c in (r, g, b))
        elif tag == "shade":
            r, g, b = (c * v for c in (r, g, b))
    return _hex([r * 255, g * 255, b * 255]), alpha


class Scheme:
    """Bảng màu theme gốc + clrMap của khuôn (để đọc đúng màu schemeClr trước khi đổi theme)."""

    def __init__(self, colors: dict, cmap: dict):
        self.colors, self.cmap = colors, cmap

    def color(self, el):
        """el: phần tử màu (srgbClr, schemeClr…). Trả về (hex, alpha) hoặc None."""
        tag = _ln(el)
        if tag == "srgbClr":
            base = el.get("val") or "000000"
        elif tag == "schemeClr":
            v = el.get("val")
            if v == "phClr":
                return None
            base = self.colors.get(self.cmap.get(v, v))
            if base is None:
                return None
        elif tag == "sysClr":
            base = el.get("lastClr") or ("000000" if el.get("val") == "windowText" else "FFFFFF")
        elif tag == "prstClr":
            base = PRST.get(el.get("val"), "000000")
        elif tag == "scrgbClr":
            base = _hex([int(el.get(k, "0")) / 100000 * 255 for k in ("r", "g", "b")])
        else:
            return None
        if len(base) != 6:
            return None
        return _mods(base, el)

    def named(self, name: str):
        """Màu theo tên vai trò (tx1, bg1…) qua clrMap."""
        c = self.colors.get(self.cmap.get(name, name))
        return (c, None) if c else None

    def of(self, parent):
        """Màu đầu tiên trong phần tử cha (solidFill, buClr, fillRef…)."""
        if parent is None:
            return None
        for c in parent:
            if _ln(c) in COLORS:
                return self.color(c)
        return None

    def fill(self, fill_el):
        """Màu đại diện của solidFill / gradFill (trung bình các điểm dừng)."""
        if fill_el is None:
            return None
        if _ln(fill_el) == "solidFill":
            return self.of(fill_el)
        if _ln(fill_el) == "gradFill":
            cs = [self.of(gs) for gs in fill_el.iter(_a("gs"))]
            cs = [c for c in cs if c]
            if not cs:
                return None
            avg = [sum(_rgb(c[0])[i] for c in cs) / len(cs) for i in range(3)]
            al = [c[1] if c[1] is not None else 1.0 for c in cs]
            return _hex(avg), sum(al) / len(al)
        return None


def _read_scheme(theme_root) -> dict:
    out = {}
    cs = theme_root.find(f".//{_a('clrScheme')}")
    for k in SCHEME_KEYS:
        el = cs.find(_a(k)) if cs is not None else None
        c = el[0] if el is not None and len(el) else None
        if c is None:
            out[k] = "000000"
        elif _ln(c) == "srgbClr":
            out[k] = c.get("val")
        else:
            out[k] = c.get("lastClr") or ("000000" if c.get("val") == "windowText" else "FFFFFF")
    return out


def _read_fonts(theme_root) -> dict:
    out = {"major": "Calibri Light", "minor": "Calibri"}
    for k, tag in (("major", "majorFont"), ("minor", "minorFont")):
        el = theme_root.find(f".//{_a(tag)}/{_a('latin')}")
        if el is not None and el.get("typeface"):
            out[k] = el.get("typeface")
    return out


def _alpha(clr) -> float:
    a = clr.find(_a("alpha"))
    return int(a.get("val")) / 100000 if a is not None else 1.0


def _srgb(hexv: str, alpha=None):
    el = etree.Element(_a("srgbClr"), val=hexv.upper())
    if alpha is not None and alpha < 1:
        etree.SubElement(el, _a("alpha"), val=str(_r(alpha * 100000)))
    return el


def _set_color(parent, hexv: str, alpha=None):
    """Thay phần tử màu trong parent (solidFill, gs, buClr…) bằng srgbClr mới, giữ độ trong suốt."""
    for c in list(parent):
        if _ln(c) in COLORS:
            parent.remove(c)
    parent.insert(0, _srgb(hexv, alpha))


def _solid(hexv: str, alpha=None):
    el = etree.Element(_a("solidFill"))
    el.append(_srgb(hexv, alpha))
    return el


# ======================================================================
# đổi màu theo mẫu
# ======================================================================

def _tdist(c: str, bg: str) -> float:
    """Màu c cách nền bg bao xa (0 = trùng nền, 1 = tương phản hết mức), tính theo độ sáng."""
    lb = lightness(bg)
    return min(1.0, abs(lightness(c) - lb) / max(lb, 1 - lb, 0.5))


class Mapper:
    def __init__(self, th: dict, order: list[int]):
        self.th, self.order = th, order

    def base(self, c: str, dark: bool) -> str:
        """Màu có sắc độ -> màu nhấn của mẫu; nhóm màu dùng nhiều nhất trong file -> màu nhấn chính."""
        th = self.th
        b = hue_bucket(c)
        k = self.order.index(b) if b in self.order else len(self.order)
        pal = [th["inv_accent"], th["accent2"], th["accent3"]] if dark else [th["accent"], th["accent2"], th["accent3"]]
        return pal[k % 3]

    def fill(self, c: str, s: "SlideCtx") -> str:
        """Màu hình khối / đường kẻ: màu xám theo khoảng cách tới nền, màu có sắc độ theo màu nhấn."""
        t = _tdist(c, s.bo)
        if chroma(c) < 0.15:
            return D.mix(s.bn, s.tn, t)
        base = self.base(c, s.inv)
        if t <= 0.65:
            return D.mix(s.bn, base, min(1.0, t * 2.2))
        return D.mix(base, s.tn, (t - 0.65) / 0.35 * 0.5)

    def text(self, c, role: str, eff_new: str, eff_old: str, size: float, bold: bool) -> str:
        th = self.th
        dark = D.contrast(eff_new, th["inv_ink"]) > D.contrast(eff_new, th["ink"])
        ink, body, muted = (th["inv_ink"], th["inv_body"], th["inv_body"]) if dark else (th["ink"], th["body"], th["muted"])
        main = ink if role == "title" else body
        colored = c is not None and chroma(c) >= 0.15
        if c is None:
            out = main
        elif not colored:
            out = main if _tdist(c, eff_old) >= 0.55 else muted
        else:
            out = self.base(c, dark)
        need = 3.0 if size >= 24 or (bold and size >= 18.5) else 4.5
        if colored:
            for _ in range(6):
                if D.contrast(eff_new, out) >= need:
                    break
                out = D.mix(out, ink, 0.3)
        if D.contrast(eff_new, out) < need:
            out = D.readable_on(eff_new, out, body, ink)
        return out


class SlideCtx:
    def __init__(self, bo, bn, tn, inv):
        self.bo, self.bn, self.tn, self.inv = bo, bn, tn, inv


# ======================================================================
# vị trí các hình (kể cả trong nhóm) theo tọa độ trang
# ======================================================================

def _xfrm_of(el):
    tag = _ln(el)
    if tag == "graphicFrame":
        return el.find(qn("p:xfrm"))
    if tag == "grpSp":
        pr = el.find(qn("p:grpSpPr"))
    else:
        pr = el.find(qn("p:spPr"))
    return pr.find(_a("xfrm")) if pr is not None else None


def _ph_of(el):
    for nv in el:
        if _ln(nv).startswith("nv"):
            ph = nv.find(f"{qn('p:nvPr')}/{qn('p:ph')}")
            if ph is not None:
                return (ph.get("type") or "body", ph.get("idx"))
    return None


def _match_ph(doc, ph):
    """Placeholder cùng loại / cùng số trong bố cục hoặc khuôn."""
    if doc is None or ph is None:
        return None
    typ, idx = ph
    norm = {"ctrTitle": "title", "subTitle": "body", "obj": "body"}
    by_type = None
    for sp in doc.iter(qn("p:sp")):
        p2 = _ph_of(sp)
        if p2 is None:
            continue
        if idx is not None and p2[1] == idx:
            return sp
        if by_type is None and norm.get(p2[0], p2[0]) == norm.get(typ, typ):
            by_type = sp
    return by_type


class Rec:
    __slots__ = ("el", "tag", "box", "z", "ph", "top", "text", "fill_old", "fill_new", "pic", "alpha")

    def __init__(self, el, box, z, ph, top):
        self.el, self.tag, self.box, self.z, self.ph, self.top = el, _ln(el), box, z, ph, top
        self.text = False
        self.fill_old = self.fill_new = None
        self.pic = False
        self.alpha = 1.0


def _walk(tree, layout_el, master_el):
    recs: list[Rec] = []

    def geo(el, ph):
        x = _xfrm_of(el)
        if (x is None or x.find(_a("off")) is None) and ph:
            for d in (layout_el, master_el):
                m = _match_ph(d, ph)
                if m is not None:
                    x = _xfrm_of(m)
                    if x is not None and x.find(_a("off")) is not None:
                        break
        if x is None or x.find(_a("off")) is None or x.find(_a("ext")) is None:
            return None
        o, e = x.find(_a("off")), x.find(_a("ext"))
        return int(o.get("x", 0)), int(o.get("y", 0)), int(e.get("cx", 0)), int(e.get("cy", 0))

    def walk(container, tf, top):
        for ch in container:
            tag = _ln(ch)
            if tag not in ("sp", "pic", "graphicFrame", "cxnSp", "grpSp"):
                if tag == "AlternateContent":   # đồ họa mới (biểu đồ 3D, SVG…): giữ nguyên, chỉ lấy vị trí
                    for sub in ch:
                        walk(sub, tf, top)
                        break
                continue
            ph = _ph_of(ch)
            g = geo(ch, ph)
            box = tf(*g) if g else (0, 0, 0, 0)
            rec = Rec(ch, box, len(recs), ph, top)
            recs.append(rec)
            if tag == "grpSp":
                x = _xfrm_of(ch)
                t2 = tf
                if x is not None and x.find(_a("off")) is not None:
                    o, e = x.find(_a("off")), x.find(_a("ext"))
                    co, ce = x.find(_a("chOff")), x.find(_a("chExt"))
                    ox, oy, ex, ey = int(o.get("x", 0)), int(o.get("y", 0)), int(e.get("cx", 0)), int(e.get("cy", 0))
                    cx = int(co.get("x", 0)) if co is not None else 0
                    cy = int(co.get("y", 0)) if co is not None else 0
                    cex = (int(ce.get("cx", 0)) if ce is not None else 0) or ex or 1
                    cey = (int(ce.get("cy", 0)) if ce is not None else 0) or ey or 1
                    sx, sy = ex / cex, ey / cey

                    def t2(X, Y, W_, H_, tf=tf, ox=ox, oy=oy, cx=cx, cy=cy, sx=sx, sy=sy):
                        return tf(ox + (X - cx) * sx, oy + (Y - cy) * sy, W_ * sx, H_ * sy)
                walk(ch, t2, False)
    walk(tree, lambda x, y, w, h: (x, y, w, h), True)
    return recs


@lru_cache(maxsize=1)
def _font_registry() -> dict:
    """Tên phông -> file, theo danh sách phông đã cài của Windows (cả phông cài riêng cho người dùng)."""
    out = {}
    try:
        import winreg
    except ImportError:
        return out
    key = r"SOFTWARE\Microsoft\Windows NT\CurrentVersion\Fonts"
    for root, base in ((winreg.HKEY_LOCAL_MACHINE, TF._DIRS[0]), (winreg.HKEY_CURRENT_USER, TF._DIRS[1])):
        try:
            k = winreg.OpenKey(root, key)
        except OSError:
            continue
        i = 0
        while True:
            try:
                name, data, _ = winreg.EnumValue(k, i)
            except OSError:
                break
            i += 1
            name = re.sub(r"\s*\((TrueType|OpenType)\)\s*$", "", name, flags=re.I)
            f = Path(str(data))
            f = f if f.is_absolute() else base / f
            for part in name.split("&"):
                out.setdefault(part.strip().lower(), f)
    return out


@lru_cache(maxsize=256)
def _orig_font(face: str, bold: bool):
    """Phông gốc để đo chữ; None nếu máy không cài phông này (khi đó giữ nguyên phông gốc, không đổi)."""
    f, k = TF._font(face, bold)
    if k == 1.0:
        return f
    reg = _font_registry()
    for name in ([f"{face} bold"] if bold else []) + [face]:
        p = reg.get(name.lower())
        if p is not None and p.exists():
            try:
                return ImageFont.truetype(str(p), TF._UNIT)
            except OSError:
                pass
    return None


def _n_lines(text: str, face: str, size: float, width_in: float, bold: bool, slack: float, spc: float,
             orig: bool = False) -> int:
    """Số dòng khi ngắt theo khoảng trắng như PowerPoint, tính cả giãn chữ spc (pt mỗi ký tự)."""
    maxw = width_in * 72 * slack
    if orig:
        f = _orig_font(face, bold)
        wid = lambda t: f.getlength(t) * size / TF._UNIT + spc * len(t)
    else:
        wid = lambda t: TF.text_width_pt(t, face, size, bold) + spc * len(t)
    n = 0
    for para in text.split("\n"):
        cur = ""
        for w in para.split(" "):
            trial = (cur + " " + w) if cur else w
            if wid(trial) <= maxw:
                cur = trial
                continue
            if cur:
                n += 1
            while wid(w) > maxw and len(w) > 4:
                cut = len(w)
                while cut > 1 and wid(w[:cut]) > maxw:
                    cut -= 1
                n += 1
                w = w[cut:]
            cur = w
        n += 1
    return n


def _contains(box, px, py):
    x, y, w, h = box
    return x <= px <= x + w and y <= py <= y + h


def _overlap(a, b) -> float:
    ax, ay, aw, ah = a
    bx, by, bw, bh = b
    return max(0, min(ax + aw, bx + bw) - max(ax, bx)) * max(0, min(ay + ah, by + bh) - max(ay, by))


# ======================================================================
# kiểu chữ thừa kế (cỡ, phông, đậm) từ hình -> bố cục -> khuôn -> trình chiếu
# ======================================================================

class Styles:
    def __init__(self, prs_el, layout_el, master_el, fonts):
        self.prs_el, self.layout_el, self.master_el, self.fonts = prs_el, layout_el, master_el, fonts

    def chain(self, txBody, ph, level: int) -> list:
        tag = _a(f"lvl{min(level, 8) + 1}pPr")
        out = []
        ls = txBody.find(_a("lstStyle")) if txBody is not None else None
        if ls is not None:
            out.append(ls.find(tag))
        tx_styles = self.master_el.find(qn("p:txStyles")) if self.master_el is not None else None
        if ph:
            for d in (self.layout_el, self.master_el):
                m = _match_ph(d, ph)
                if m is not None:
                    ls = m.find(f"{qn('p:txBody')}/{_a('lstStyle')}")
                    if ls is not None:
                        out.append(ls.find(tag))
            if tx_styles is not None:
                key = "p:titleStyle" if ph[0] in ("title", "ctrTitle") else "p:bodyStyle"
                st = tx_styles.find(qn(key))
                if st is not None:
                    out.append(st.find(tag))
        else:
            dts = self.prs_el.find(qn("p:defaultTextStyle"))
            if dts is not None:
                out.append(dts.find(tag))
            if tx_styles is not None and tx_styles.find(qn("p:otherStyle")) is not None:
                out.append(tx_styles.find(qn("p:otherStyle")).find(tag))
        return [e for e in out if e is not None]

    @staticmethod
    def attr(rPr, chain, name):
        if rPr is not None and rPr.get(name) is not None:
            return rPr.get(name)
        for lv in chain:
            d = lv.find(_a("defRPr"))
            if d is not None and d.get(name) is not None:
                return d.get(name)
        return None

    def face(self, rPr, chain, ph):
        el = rPr.find(_a("latin")) if rPr is not None else None
        face = el.get("typeface") if el is not None else None
        if not face:
            for lv in chain:
                d = lv.find(f"{_a('defRPr')}/{_a('latin')}")
                if d is not None and d.get("typeface"):
                    face = d.get("typeface")
                    break
        if not face:
            face = "+mj-lt" if ph and ph[0] in ("title", "ctrTitle") else "+mn-lt"
        if face.startswith("+mj"):
            return self.fonts["major"]
        if face.startswith("+mn"):
            return self.fonts["minor"]
        return face


# ======================================================================
# xử lý từng trang
# ======================================================================

RPR_ORDER = ("ln", "noFill", "solidFill", "gradFill", "blipFill", "pattFill", "grpFill", "effectLst", "effectDag",
             "highlight", "uLnTx", "uLn", "uFillTx", "uFill", "latin", "ea", "cs", "sym", "hlinkClick",
             "hlinkMouseOver", "rtl", "extLst")
SPPR_ORDER = ("xfrm", "custGeom", "prstGeom") + FILLS + ("ln", "effectLst", "effectDag", "scene3d", "sp3d", "extLst")
TCPR_ORDER = ("lnL", "lnR", "lnT", "lnB", "lnTlToBr", "lnBlToTr", "cell3D") + FILLS + ("headers", "extLst")


def _put(parent, el, order):
    """Chèn el vào parent đúng thứ tự phần tử con theo lược đồ OOXML (thay phần tử cùng loại nếu có)."""
    name = _ln(el)
    group = FILLS if name in FILLS else (name,)
    for c in list(parent):
        if _ln(c) in group:
            parent.remove(c)
    rank = order.index(name)
    for i, c in enumerate(parent):
        n = _ln(c)
        if n in order and order.index(n) > rank:
            parent.insert(i, el)
            return el
    parent.append(el)
    return el


def _fill_el(pr):
    if pr is None:
        return None
    for c in pr:
        if _ln(c) in FILLS:
            return c
    return None


def _txt(txBody) -> str:
    return "".join(t.text or "" for t in txBody.iter(_a("t"))) if txBody is not None else ""


class Restyler:
    def __init__(self, prs, th: dict, opts: dict):
        self.prs, self.th, self.opts = prs, th, opts
        self.SW, self.SH = prs.slide_width or 12192000, prs.slide_height or 6858000
        self.themes = {}   # theme part -> (scheme, fonts) gốc
        for m in prs.slide_masters:
            part = m.part.part_related_by(RT.THEME)
            if part not in self.themes:
                root = etree.fromstring(part.blob)
                self.themes[part] = (_read_scheme(root), _read_fonts(root))
        self.mapper = Mapper(th, self._bucket_order())

    # ---------- màu gốc của từng trang
    def _scheme(self, slide) -> Scheme:
        master = slide.slide_layout.slide_master
        colors, _ = self.themes[master.part.part_related_by(RT.THEME)]
        cm = master._element.find(qn("p:clrMap"))
        cmap = dict(cm.attrib) if cm is not None else {}
        ov = slide._element.find(f"{qn('p:clrMapOvr')}/{_a('overrideClrMapping')}")
        if ov is not None:
            cmap = dict(ov.attrib)
        return Scheme(colors, cmap)

    def _bucket_order(self) -> list[int]:
        """Đếm các nhóm màu có sắc độ trong cả file để biết nhóm nào là màu chủ đạo."""
        cnt: dict[int, int] = {}
        first: dict[int, int] = {}
        n = 0
        for slide in self.prs.slides:
            sc = self._scheme(slide)
            for el in slide._element.iter(_a("solidFill"), _a("gs")):
                c = sc.of(el)
                if c and chroma(c[0]) >= 0.15:
                    b = hue_bucket(c[0])
                    cnt[b] = cnt.get(b, 0) + 1
                    first.setdefault(b, n)
                    n += 1
        return sorted(cnt, key=lambda b: (-cnt[b], first[b]))

    def _orig_bg(self, slide, sc: Scheme) -> str:
        for el in (slide._element, slide.slide_layout._element, slide.slide_layout.slide_master._element):
            bg = el.find(f"{qn('p:cSld')}/{qn('p:bg')}")
            if bg is None:
                continue
            pr = bg.find(qn("p:bgPr"))
            if pr is not None:
                c = sc.fill(_fill_el(pr))
                return c[0] if c else "FFFFFF"
            ref = bg.find(qn("p:bgRef"))
            if ref is not None:
                c = sc.of(ref)
                return c[0] if c else "FFFFFF"
        return "FFFFFF"

    # ---------- cả bộ
    def run(self, kinds: list[str]):
        for i, slide in enumerate(self.prs.slides):
            self.slide(slide, i, kinds[i] if i < len(kinds) else "content")
        for part in self.themes:
            self._restyle_theme(part)
        self.prs.core_properties.subject = f"Giữ bố cục gốc – mẫu {self.th['name']}"

    def _restyle_theme(self, part):
        th = self.th
        root = etree.fromstring(part.blob)
        vals = dict(dk1=th["ink"], lt1=th["bg"], dk2=th["inv"][0], lt2=th["bg2"], accent1=th["accent"],
                    accent2=th["accent2"], accent3=th["accent3"], accent4=D.mix(th["accent"], th["accent2"], 0.5),
                    accent5=D.mix(th["accent2"], th["accent3"], 0.5), accent6=D.mix(th["accent"], th["ink"], 0.35),
                    hlink=th["accent"], folHlink=th["muted"])
        cs = root.find(f".//{_a('clrScheme')}")
        if cs is not None:
            for k, v in vals.items():
                el = cs.find(_a(k))
                if el is not None:
                    for c in list(el):
                        el.remove(c)
                    el.append(_srgb(v))
        for tag, face in (("majorFont", th["head"]), ("minorFont", th["body_font"])):
            el = root.find(f".//{_a(tag)}/{_a('latin')}")
            if el is not None:
                el.set("typeface", face)
                for k in ("panose", "pitchFamily", "charset"):
                    el.attrib.pop(k, None)
        part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    # ---------- một trang
    def slide(self, slide, idx: int, kind: str):
        th, opts = self.th, self.opts
        sld = slide._element
        tree = sld.find(qn("p:cSld")).find(qn("p:spTree"))
        sc = self._scheme(slide)
        bo = self._orig_bg(slide, sc)
        inv = kind in ("cover", "section", "closing")
        styles = Styles(self.prs._element, slide.slide_layout._element, slide.slide_layout.slide_master._element,
                        self.themes[slide.slide_layout.slide_master.part.part_related_by(RT.THEME)][1])
        had_timing = sld.find(qn("p:timing")) is not None
        recs = _walk(tree, slide.slide_layout._element, slide.slide_layout.slide_master._element)

        # 1. nền + trang trí của mẫu (vẽ xong thì dời xuống lớp dưới cùng)
        before = list(tree)
        info = decor.paint(slide, th, inv, idx, kind if inv else "content")
        added = [c for c in tree if c not in before]
        self._place_decor(tree, added)
        bn = info["bg"]
        tn = D.readable_on(bn, th["inv_ink"] if inv else th["ink"], th["ink"], th["inv_ink"])
        s = SlideCtx(bo, bn, tn, inv)
        sld.set("showMasterSp", "0")   # ẩn hình trang trí của mẫu cũ (trong slide master)
        self._clrmap(sld, bn)

        # 2. màu hình khối, đường kẻ
        for r in recs:
            self._shape_colors(r, s, sc)
        # 3. chữ (trong hình và trong bảng)
        for r in recs:
            if r.tag == "sp":
                tx = r.el.find(qn("p:txBody"))
                if tx is not None and _txt(tx).strip():
                    r.text = True
                    self._text_body(tx, r, recs, s, sc, styles, r.box[2])
            elif r.tag == "graphicFrame" and r.el.find(f".//{_a('tbl')}") is not None:
                r.text = True
                self._table(r, recs, s, sc, styles)
            elif r.tag == "graphicFrame":
                r.text = True   # biểu đồ, SmartArt: không đổi gì, chỉ tránh để họa tiết đè lên
        # 4. họa tiết chạm vào chữ, ảnh thì làm mờ (ảnh phủ kín trang che hết họa tiết nên bỏ qua)
        full = 0.85 * self.SW * self.SH
        self._soften_decor(added, [r.box for r in recs if r.text or (r.pic and r.box[2] * r.box[3] < full)])
        # 5. animation (giữ animation có sẵn) + chuyển trang
        if not had_timing and opts.get("anim", "auto") != "none":
            add_animations(slide, self._steps(slide, recs, added), opts.get("anim", "auto"))
        add_transition(slide, opts.get("transition") or th["transition"])

    def _clrmap(self, sld, bn):
        """Trang nền tối: đảo clrMap để chữ thừa kế màu theme (biểu đồ, SmartArt) vẫn đọc được."""
        th = self.th
        swap = D.contrast(bn, th["ink"]) < 3 and D.contrast(bn, th["bg"]) > D.contrast(bn, th["ink"])
        ovr = sld.find(qn("p:clrMapOvr"))
        if ovr is None:
            ovr = etree.SubElement(sld, qn("p:clrMapOvr"))
            sld.find(qn("p:cSld")).addnext(ovr)
        for c in list(ovr):
            ovr.remove(c)
        if swap:
            m = dict(bg1="dk1", tx1="lt1", bg2="dk2", tx2="lt2", accent1="accent1", accent2="accent2",
                     accent3="accent3", accent4="accent4", accent5="accent5", accent6="accent6", hlink="hlink",
                     folHlink="folHlink")
            etree.SubElement(ovr, _a("overrideClrMapping"), **m)
        else:
            etree.SubElement(ovr, _a("masterClrMapping"))

    # ---------- trang trí: co giãn theo khổ trang, đặt dưới cùng
    def _place_decor(self, tree, added):
        sx, sy = self.SW / D.E(D.W), self.SH / D.E(D.H)
        if abs(sx - 1) > 1e-6 or abs(sy - 1) > 1e-6:
            s = min(sx, sy)
            for el in added:
                x = _xfrm_of(el)
                if x is None or x.find(_a("off")) is None:
                    continue
                o, e = x.find(_a("off")), x.find(_a("ext"))
                X, Y, Wd, Hd = (int(o.get("x")), int(o.get("y")), int(e.get("cx")), int(e.get("cy")))
                if Wd > 0.6 * D.E(D.W) or Hd > 0.6 * D.E(D.H):     # hình trải ngang/dọc trang: kéo theo khổ
                    X, Y, Wd, Hd = X * sx, Y * sy, Wd * sx, Hd * sy
                else:                                            # hình gọn: giữ tỉ lệ, dời tâm theo khổ
                    cx, cy = (X + Wd / 2) * sx, (Y + Hd / 2) * sy
                    Wd, Hd = Wd * s, Hd * s
                    X, Y = cx - Wd / 2, cy - Hd / 2
                o.set("x", str(_r(X)))
                o.set("y", str(_r(Y)))
                e.set("cx", str(max(1, _r(Wd))))
                e.set("cy", str(max(1, _r(Hd))))
        pos = list(tree).index(tree.find(qn("p:grpSpPr"))) + 1
        for k, el in enumerate(added):
            tree.remove(el)
            tree.insert(pos + k, el)

    def _soften_decor(self, added, boxes):
        pad = int(0.04 * EMU)
        for el in added:
            x = _xfrm_of(el)
            if x is None or x.find(_a("off")) is None:
                continue
            o, e = x.find(_a("off")), x.find(_a("ext"))
            box = (int(o.get("x")), int(o.get("y")), int(e.get("cx")), int(e.get("cy")))
            spPr = el.find(qn("p:spPr"))
            f = _fill_el(spPr)
            strong = f is not None and _ln(f) in ("solidFill", "gradFill") and any(
                _alpha(c) > 0.3 for c in f.iter(_a("srgbClr")))
            tx = el.find(qn("p:txBody"))
            has_text = tx is not None and _txt(tx).strip()
            inner = (box[0] + pad * 2, box[1] + pad * 2, box[2] - pad * 4, box[3] - pad * 4)
            hit = False
            for b in boxes:
                if _overlap(box, b) <= 0.02 * EMU * EMU:
                    continue
                if strong or has_text:
                    hit = True
                elif not (inner[0] <= b[0] and inner[1] <= b[1] and b[0] + b[2] <= inner[0] + inner[2]
                          and b[1] + b[3] <= inner[1] + inner[3]):
                    hit = True   # khung viền: chỉ tính khi chữ cắt qua đường viền
                if hit:
                    break
            if not hit:
                continue
            if has_text:
                el.getparent().remove(el)
                continue
            for c in spPr.iter(_a("srgbClr")):
                ln = next((p for p in c.iterancestors() if _ln(p) == "ln"), None)
                cap = min(_alpha(c), 0.3 if ln is not None else 0.18)
                a = c.find(_a("alpha"))
                if a is None:
                    a = etree.SubElement(c, _a("alpha"))
                a.set("val", str(_r(cap * 100000)))

    # ---------- màu hình
    def _shape_colors(self, r: Rec, s: SlideCtx, sc: Scheme):
        el, M = r.el, self.mapper
        if r.tag == "pic":
            r.pic = True
            self._line(el.find(qn("p:spPr")), el, s, sc)
            return
        if r.tag not in ("sp", "cxnSp", "grpSp"):
            return
        spPr = el.find(qn("p:spPr") if r.tag != "grpSp" else qn("p:grpSpPr"))
        if spPr is None:
            return
        f = _fill_el(spPr)
        style = el.find(qn("p:style"))
        if f is None and style is not None and r.tag == "sp":
            ref = style.find(_a("fillRef"))
            if ref is not None and ref.get("idx", "0") not in ("0", "1000"):
                c = sc.of(ref)
                if c:
                    f = _put(spPr, _solid(c[0], c[1]), SPPR_ORDER)
        if f is not None:
            kind = _ln(f)
            if kind == "blipFill":
                r.pic = True
            elif kind == "solidFill":
                c = sc.of(f)
                if c:
                    if self._is_backdrop(r, spPr):
                        _put(spPr, etree.Element(_a("noFill")), SPPR_ORDER)
                    else:
                        new = M.fill(c[0], s)
                        _set_color(f, new, c[1])
                        r.fill_old, r.fill_new, r.alpha = c[0], new, (c[1] if c[1] is not None else 1.0)
            elif kind == "gradFill":
                avg = sc.fill(f)
                if avg and self._is_backdrop(r, spPr):
                    _put(spPr, etree.Element(_a("noFill")), SPPR_ORDER)
                else:
                    news = []
                    for gs in f.iter(_a("gs")):
                        c = sc.of(gs)
                        if c:
                            new = M.fill(c[0], s)
                            _set_color(gs, new, c[1])
                            news.append(new)
                    if avg and news:
                        r.fill_old = avg[0]
                        r.fill_new = _hex([sum(_rgb(n)[i] for n in news) / len(news) for i in range(3)])
                        r.alpha = avg[1]
        if r.tag != "grpSp":
            self._line(spPr, el, s, sc)

    def _is_backdrop(self, r: Rec, spPr) -> bool:
        """Mảng màu phủ gần kín trang, không chữ: là nền của mẫu cũ => bỏ để lộ nền mới."""
        geom = spPr.find(_a("prstGeom"))
        if r.tag != "sp" or geom is None or geom.get("prst") != "rect" or _txt(r.el.find(qn("p:txBody"))).strip():
            return False
        x, y, w, h = r.box
        cover = max(0, min(x + w, self.SW) - max(x, 0)) * max(0, min(y + h, self.SH) - max(y, 0))
        return cover >= 0.92 * self.SW * self.SH

    def _line(self, spPr, el, s, sc):
        if spPr is None:
            return
        ln = spPr.find(_a("ln"))
        f = _fill_el(ln) if ln is not None else None
        if f is not None:
            if _ln(f) == "solidFill":
                c = sc.of(f)
                if c:
                    _set_color(f, self.mapper.fill(c[0], s), c[1])
            elif _ln(f) == "gradFill":
                for gs in f.iter(_a("gs")):
                    c = sc.of(gs)
                    if c:
                        _set_color(gs, self.mapper.fill(c[0], s), c[1])
            return
        style = el.find(qn("p:style"))
        ref = style.find(_a("lnRef")) if style is not None else None
        if ref is not None and ref.get("idx", "0") != "0":
            c = sc.of(ref)
            if c:
                if ln is None:
                    ln = _put(spPr, etree.Element(_a("ln")), SPPR_ORDER)
                ln.insert(0, _solid(self.mapper.fill(c[0], s), c[1]))

    # ---------- chữ
    def _eff_bg(self, r: Rec, recs, s: SlideCtx):
        """Nền ngay dưới chữ: (màu gốc, màu mới) hoặc None nếu chữ nằm trên ảnh (giữ màu chữ gốc)."""
        if r.pic:
            return None
        if r.fill_old and r.alpha >= 0.5:
            return r.fill_old, r.fill_new
        x, y, w, h = r.box
        cx, cy = x + w / 2, y + h / 2
        for q in reversed(recs[:r.z]):
            if q.tag == "grpSp" or not _contains(q.box, cx, cy):
                continue
            if q.pic:
                return None
            if q.fill_old and q.alpha >= 0.5:
                return q.fill_old, q.fill_new
        return s.bo, s.bn

    def _text_body(self, tx, r: Rec, recs, s, sc, styles, width_emu, eff=None, cell=False):
        """width_emu: bề rộng hộp; trong bảng (cell) là bề rộng ô đã trừ lề ô, eff là nền ô (None: ô trên ảnh)."""
        th, M = self.th, self.mapper
        eff = eff if cell else self._eff_bg(r, recs, s)
        keep_color = eff is None
        bp = tx.find(_a("bodyPr"))
        if bp is not None and (bp.get("wrap") == "none" or bp.get("vert") not in (None, "horz")):
            width_emu = 0   # chữ không xuống dòng hoặc chữ dọc: không đo
        elif not cell:
            width_emu -= int(bp.get("lIns", "91440") if bp is not None else 91440)
            width_emu -= int(bp.get("rIns", "91440") if bp is not None else 91440)
        scale = 1.0
        na = bp.find(_a("normAutofit")) if bp is not None else None
        if na is not None and na.get("fontScale"):
            scale = int(na.get("fontScale")) / 100000
        title_ph = r.ph is not None and r.ph[0] in ("title", "ctrTitle")
        st = r.el.find(qn("p:style")) if r.tag == "sp" else None
        base = sc.of(st.find(_a("fontRef"))) if st is not None and st.find(_a("fontRef")) is not None else None
        plan = []   # (đoạn, cỡ, phông gốc, phông mới, đậm, lùi đầu dòng, chữ)
        for p in tx.findall(_a("p")):
            pPr = p.find(_a("pPr"))
            level = int(pPr.get("lvl", "0")) if pPr is not None else 0
            chain = styles.chain(tx, r.ph, level)
            runs = [c for c in p if _ln(c) in ("r", "fld")]
            text = "".join((c.findtext(_a("t")) or "") if _ln(c) != "br" else "\n" for c in p if _ln(c) in ("r", "fld", "br"))
            sizes, faces, bolds = [], [], []
            for rn in runs:
                rPr = rn.find(_a("rPr"))
                sz = Styles.attr(rPr, chain, "sz")
                sizes.append(int(sz) / 100 if sz else 18.0)
                faces.append(styles.face(rPr, chain, r.ph))
                bolds.append(Styles.attr(rPr, chain, "b") in ("1", "true"))
            size = max(sizes) if sizes else 18.0
            role = "title" if title_ph or size * scale >= 28 else "body"
            font = th["head"] if role == "title" else th["body_font"]
            if not keep_color:
                bc = pPr.find(_a("buClr")) if pPr is not None else None
                if bc is not None and sc.of(bc):
                    _set_color(bc, M.text(sc.of(bc)[0], "body", eff[1], eff[0], size * scale, False))
            # phông gốc máy không có (không đo được) hoặc phông ký hiệu: giữ nguyên để không vỡ bố cục
            keep_face = [bool(SYMBOL_FONT_RE.match(f or "")) or _orig_font(f, b) is None for f, b in zip(faces, bolds)]
            for rn, sz, face, b, kf in zip(runs, sizes, faces, bolds, keep_face):
                rPr = rn.find(_a("rPr"))
                if rPr is None:
                    rPr = etree.Element(_a("rPr"))
                    rn.insert(0, rPr)
                self._run(rPr, sc, role, eff, keep_color, sz * scale, b, None if kf else font, chain, base)
            end = p.find(_a("endParaRPr"))
            if end is not None:
                f0 = faces[0] if faces else styles.face(end, chain, r.ph)
                kf = keep_face[0] if keep_face else (bool(SYMBOL_FONT_RE.match(f0 or "")) or _orig_font(f0, False) is None)
                self._run(end, sc, role, eff, keep_color, size * scale, bolds[0] if bolds else False,
                          None if kf else font, chain, base)
            changed = [k for k, kf in enumerate(keep_face) if not kf]
            if text.strip() and changed:
                marl = pPr.get("marL") if pPr is not None and pPr.get("marL") else next(
                    (lv.get("marL") for lv in chain if lv.get("marL")), "0")
                r0 = runs[0].find(_a("rPr")) if runs else None
                if Styles.attr(r0, chain, "cap") in ("all", "small"):
                    text = text.upper()
                spc = int(Styles.attr(r0, chain, "spc") or 0) / 100   # giãn chữ (pt mỗi ký tự)
                plan.append((p, size, faces[changed[0]], font, any(bolds), int(marl), text, runs, end,
                             sizes, spc))
        self._keep_lines(plan, width_emu / EMU, scale)

    def _run(self, rPr, sc, role, eff, keep_color, size, bold, font, chain, base):
        """Tô màu một run theo mẫu; font=None: giữ phông gốc."""
        if keep_color:
            # chữ nằm trên ảnh: giữ đúng màu gốc (ghi rõ ra để không đổi theo bảng màu theme mới)
            f = _fill_el(rPr)
            if f is None:
                c = next((sc.of(lv.find(f"{_a('defRPr')}/{_a('solidFill')}")) for lv in chain
                          if lv.find(f"{_a('defRPr')}/{_a('solidFill')}") is not None), None) or base or sc.named("tx1")
                if c:
                    _put(rPr, _solid(c[0], c[1]), RPR_ORDER)
            elif _ln(f) == "solidFill" and sc.of(f):
                c = sc.of(f)
                _set_color(f, c[0], c[1])
        else:
            f = _fill_el(rPr)
            c = sc.fill(f) if f is not None and _ln(f) in ("solidFill", "gradFill") else None
            if not (f is not None and _ln(f) == "noFill"):
                new = self.mapper.text(c[0] if c else None, role, eff[1], eff[0], size, bold)
                _put(rPr, _solid(new, c[1] if c else None), RPR_ORDER)
        if font:
            lat = rPr.find(_a("latin"))
            if lat is None:
                lat = _put(rPr, etree.Element(_a("latin")), RPR_ORDER)
            lat.set("typeface", font)
            for k in ("panose", "pitchFamily", "charset"):
                lat.attrib.pop(k, None)
            cs = rPr.find(_a("cs"))
            if cs is not None:
                cs.set("typeface", font)

    def _keep_lines(self, plan, width, scale):
        """Phông mới rộng hơn làm chữ xuống thêm dòng => thu nhỏ chữ vừa đủ để giữ số dòng như bản gốc."""
        if not plan or width <= 0.3:
            return

        def lines(use_new: bool, f: float) -> int:
            # bản gốc: đo đúng bề rộng (PowerPoint đã ngắt dòng như vậy); bản mới: chừa 1,5% cho chắc
            n = 0
            for p, size, face0, face1, bold, marl, text, runs, end, sizes, spc in plan:
                w = max(0.3, width - marl / EMU)
                if use_new:
                    n += _n_lines(text, face1, size * scale * f, w, bold, 0.985, spc)
                else:
                    n += _n_lines(text, face0, size * scale, w, bold, 1.0, spc, orig=True)
            return n
        base = lines(False, 1)
        if lines(True, 1) <= base:
            return
        pick = 0.7
        for f in (0.95, 0.9, 0.85, 0.8, 0.75, 0.7):
            if lines(True, f) <= base:
                pick = f
                break
        for p, size, face0, face1, bold, marl, text, runs, end, sizes, spc in plan:
            for rn, sz in zip(runs, sizes):
                rn.find(_a("rPr")).set("sz", str(max(100, _r(sz * pick * 100))))
            if end is not None:
                end.set("sz", str(max(100, _r(size * pick * 100))))

    # ---------- bảng
    def _table(self, r: Rec, recs, s, sc, styles):
        th, M = self.th, self.mapper
        under = self._eff_bg(r, recs, s)   # None: bảng nằm trên ảnh => ô không tô màu giữ nguyên màu chữ, màu viền
        tbl = r.el.find(f".//{_a('tbl')}")
        pr = tbl.find(_a("tblPr"))
        styled = pr is not None and pr.find(_a("tableStyleId")) is not None
        first = pr is not None and pr.get("firstRow") in ("1", "true")
        band = pr is not None and pr.get("bandRow") in ("1", "true")
        cols = [int(g.get("w", "0")) for g in tbl.iter(_a("gridCol"))]
        head = th["inv_accent"] if s.inv else th["accent"]
        for i, tr in enumerate(tbl.findall(_a("tr"))):
            ci = 0
            for tc in tr.findall(_a("tc")):
                span = int(tc.get("gridSpan", "1"))
                w = sum(cols[ci:ci + span]) or (cols[ci] if ci < len(cols) else r.box[2])
                ci += span
                if tc.get("hMerge") or tc.get("vMerge"):
                    continue
                tcPr = tc.find(_a("tcPr"))
                if tcPr is None:
                    tcPr = etree.Element(_a("tcPr"))
                    tx = tc.find(_a("txBody"))
                    (tx.addnext(tcPr) if tx is not None else tc.insert(0, tcPr))
                f = _fill_el(tcPr)
                c = sc.fill(f) if f is not None else None
                if c:
                    old, new = c[0], M.fill(c[0], s)
                    _put(tcPr, _solid(new, c[1]), TCPR_ORDER)
                elif f is not None and _ln(f) == "noFill" or not styled:
                    old, new = under if under else (None, None)
                else:   # ô theo kiểu bảng của theme: tô màu rõ ràng như bảng của công cụ
                    if first and i == 0:
                        old, new = "000000", head
                    else:
                        new = th["card"] if not band or i % 2 else D.mix(th["accent"], th["card"], 0.92)
                        old = "FFFFFF"
                    _put(tcPr, _solid(new), TCPR_ORDER)
                for side in ("lnL", "lnR", "lnT", "lnB") if old else ():
                    ln = tcPr.find(_a(side))
                    lf = _fill_el(ln) if ln is not None else None
                    if lf is not None and _ln(lf) == "solidFill" and sc.of(lf):
                        cc = sc.of(lf)
                        _set_color(lf, M.fill(cc[0], s), cc[1])
                tx = tc.find(_a("txBody"))
                if tx is not None and _txt(tx).strip():
                    mar = int(tcPr.get("marL", "91440")) + int(tcPr.get("marR", "91440"))
                    cell = Rec(r.el, r.box, r.z, None, False)
                    self._text_body(tx, cell, [], s, sc, styles, w - mar, eff=(old, new) if old else None, cell=True)

    # ---------- animation: các khối nội dung theo thứ tự đọc (trên xuống, trái sang)
    def _steps(self, slide, recs, added):
        th, opts = self.th, self.opts
        top = {r.el: r for r in recs if r.top}
        cands = []
        for shp in slide.shapes:
            r = top.get(shp._element)
            if r is None or shp._element in added:
                continue
            if r.ph and r.ph[0] in ("sldNum", "dt", "ftr"):
                continue
            x, y, w, h = r.box
            if w <= 0 or h <= 0 or x >= self.SW or y >= self.SH or x + w <= 0 or y + h <= 0:
                continue
            has = r.text or r.tag in ("pic", "graphicFrame") or (
                r.tag == "grpSp" and (r.el.find(f".//{_a('t')}") is not None or r.el.find(f".//{qn('p:pic')}") is not None))
            if not has or (r.tag == "sp" and not r.text):
                continue
            if w * h >= 0.85 * self.SW * self.SH:   # ảnh/khối phủ kín trang: là nền, không cho bay vào
                continue
            cands.append((r, shp))
        if not cands:
            return []
        cands.sort(key=lambda t: (t[0].box[1], t[0].box[0]))
        title = next((t for t in cands if t[0].ph and t[0].ph[0] in ("title", "ctrTitle")), None)
        if title:
            cands.remove(title)
        eff = opts.get("effect") or th["effect"]
        steps, band = [], None
        for r, shp in cands:
            if steps and r.box[1] < band + 0.06 * self.SH:
                steps[-1].append((shp, eff))
            else:
                steps.append([(shp, eff)])
                band = r.box[1]
        while len(steps) > 7:
            steps[-2].extend(steps.pop())
        if title:
            steps.insert(0, [(title[1], "fade")])
        return steps


# ======================================================================
# loại trang: bìa, chương, kết (nền đậm) hay nội dung
# ======================================================================

def slide_kinds(path, prs, restyler: Restyler) -> list[str]:
    deck = readers.read_pptx(Path(path))
    raw = [Slide(title=deck.title, subtitle=deck.subtitle)] + list(deck.slides)
    n = len(prs.slides)
    out = []
    for i, sl in enumerate(prs.slides):
        s = raw[i] if i < len(raw) else Slide()
        bg = restyler._orig_bg(sl, restyler._scheme(sl))
        dark = lightness(bg) < 0.45 or (chroma(bg) >= 0.3 and lightness(bg) < 0.75)   # nền tối hoặc nền màu đậm
        texts = [s.title] + [it.text for it in s.items]
        if i == 0:
            out.append("cover")
        elif i == n - 1 and any(CLOSING_RE.search(t or "") and len(t) < 80 for t in texts):
            out.append("closing")
        elif dark or (s.title and not s.items and not s.images and not s.table and i < n - 1):
            out.append("section")
        else:
            out.append("content")
    return out


def build(src, th: dict, out_dir, opts: dict) -> Path:
    prs = Presentation(str(src))
    rs = Restyler(prs, th, opts)
    rs.run(slide_kinds(src, prs, rs))
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    stem = re.sub(r'[\\/:*?"<>|]+', "_", Path(src).stem)[:80]
    out = out_dir / f"{stem} - {th['name']}.pptx"
    k = 2
    while out.exists():
        try:
            out.unlink()  # ghi đè bản cũ nếu không bị mở trong PowerPoint
            break
        except PermissionError:
            out = out_dir / f"{stem} - {th['name']} ({k}).pptx"
            k += 1
    prs.save(str(out))
    return out
