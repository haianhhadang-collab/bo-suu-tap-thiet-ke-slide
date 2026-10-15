"""Dựng bộ slide hoàn chỉnh từ một file đầu vào."""
from __future__ import annotations

import re
from pathlib import Path

from pptx import Presentation

from . import layouts as L
from . import hoavan, readers, themes
from .draw import E, H, W
from .model import Deck, Item, Slide


def _is_section(s: Slide) -> bool:
    return s.hint == "section" or (bool(s.title) and not s.items and not s.images and not s.table)


def _split_closing(deck: Deck):
    """Nếu trang cuối có câu cảm ơn lẫn trong nội dung thì tách thành trang kết riêng."""
    if not deck.slides:
        return
    last = deck.slides[-1]
    if L.CLOSING_RE.search(last.title or "") and len(last.items) <= 4 and not last.images and not last.table:
        last.hint = "closing"
        return
    hits = [i for i in last.items if L.CLOSING_RE.search(i.text) and len(i.text) < 80]
    if hits and not last.title and len(last.items) == len(hits):
        last.title, last.items, last.hint = hits[0].text, hits[1:], "closing"
        return
    if hits and len(last.items) > len(hits):
        last.items = [i for i in last.items if i not in hits]
        deck.slides.append(Slide(title=hits[0].text, hint="closing",
                                 items=[Item(text=h.text, kind="para") for h in hits[1:]]))


def build(src, theme_id: str, out_dir, opts: dict | None = None, deck: Deck | None = None) -> Path:
    opts = dict(opts or {})
    # "co-dien+dong-son" = mẫu co-dien kèm họa tiết dong-son; không ghi "+" thì lấy họa tiết trong opts
    theme_id, _, motif = theme_id.partition("+")
    th = hoavan.apply(themes.get(theme_id), motif or opts.get("motif", ""), opts.get("motif_palette", True))
    deck = deck or readers.read(src)
    _split_closing(deck)

    prs = Presentation()
    prs.slide_width, prs.slide_height = E(W), E(H)
    L.cover(prs, th, deck, opts)

    queue = list(deck.slides)
    idx, page, secnum, guard = 1, 2, 0, 0
    while queue and guard < 400:
        guard += 1
        s = queue.pop(0)
        if s.hint == "closing" or (not queue and _is_section(s) and L.CLOSING_RE.search(s.title or "")):
            L.closing(prs, th, s, idx, opts)
        elif _is_section(s) and queue:
            secnum += 1
            L.section(prs, th, s, idx, secnum, opts)
        else:
            rest = L.content(prs, th, s, idx, page, opts)
            if rest is not None and (rest.items or rest.images or rest.table):
                if rest.title and not rest.title.endswith("(tiếp)"):
                    rest.title = rest.title + " (tiếp)"
                queue.insert(0, rest)
        idx += 1
        page += 1

    prs.core_properties.title = deck.title[:250]
    prs.core_properties.subject = f"Dàn trang tự động – mẫu {th['name']}"
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    stem = re.sub(r'[\\/:*?"<>|]+', "_", Path(src).stem)[:80]
    out = out_dir / f"{stem} - {th['name']}.pptx"
    k = 2
    while out.exists():
        try:
            out.unlink()  # ghi đè bản cũ nếu không bị mở trong PowerPoint
            break
        except PermissionError:
            out = out_dir / f"{stem} - {th['name']} ({k}).pptx"
            k += 1
    prs.save(str(out))
    return out


def build_all(src, out_dir, opts=None) -> list[Path]:
    deck_once = readers.read(src)
    outs = []
    import copy
    for tid in themes.THEMES:
        outs.append(build(src, tid, out_dir, opts, deck=copy.deepcopy(deck_once)))
    return outs
