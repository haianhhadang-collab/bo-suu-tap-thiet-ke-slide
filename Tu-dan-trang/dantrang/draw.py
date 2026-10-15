"""Các hàm vẽ trên slide python-pptx: khối, chữ, ảnh, nền gradient, độ trong suốt, bóng đổ."""
from __future__ import annotations

import io
import math

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

EMU = 914400
W, H = 13.333, 7.5
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def E(x: float) -> Emu:
    return Emu(int(round(x * EMU)))


def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h.lstrip("#").upper())


def mix(a: str, b: str, t: float) -> str:
    """Trộn màu a→b theo tỉ lệ t (0..1)."""
    a, b = a.lstrip("#"), b.lstrip("#")
    ca = [int(a[i:i + 2], 16) for i in (0, 2, 4)]
    cb = [int(b[i:i + 2], 16) for i in (0, 2, 4)]
    # làm tròn 0.5 lên như Math.round của bản web (round() của Python làm tròn về số chẵn)
    return "".join(f"{math.floor(x + (y - x) * t + 0.5):02X}" for x, y in zip(ca, cb))


def luminance(h: str) -> float:
    h = h.lstrip("#")
    r, g, b = (int(h[i:i + 2], 16) / 255 for i in (0, 2, 4))
    f = lambda c: c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4
    return 0.2126 * f(r) + 0.7152 * f(g) + 0.0722 * f(b)


def contrast(a: str, b: str) -> float:
    la, lb = sorted((luminance(a), luminance(b)), reverse=True)
    return (la + 0.05) / (lb + 0.05)


def readable_on(bg: str, *cands: str) -> str:
    """Chọn màu chữ đầu tiên đạt 4.5:1 trên nền bg, nếu không thì màu tương phản nhất."""
    for c in cands:
        if contrast(bg, c) >= 4.5:
            return c
    return max(cands + ("FFFFFF", "111111"), key=lambda c: contrast(bg, c))


# ---------- tô màu (đặc, gradient, trong suốt) ----------

def _clr(hexv: str, alpha: float | None = None) -> str:
    a = f'<a:alpha val="{int(alpha * 100000)}"/>' if alpha is not None and alpha < 1 else ""
    return f'<a:srgbClr val="{hexv.lstrip("#").upper()}">{a}</a:srgbClr>'


def grad_xml(colors, angle: float = 90, radial: bool = False, alphas=None, positions=None) -> str:
    n = len(colors)
    pos = positions or [round(i * 100000 / max(n - 1, 1)) for i in range(n)]
    al = alphas or [None] * n
    gs = "".join(f'<a:gs pos="{p}">{_clr(c, a)}</a:gs>' for c, p, a in zip(colors, pos, al))
    if radial:
        shade = '<a:path path="circle"><a:fillToRect l="50000" t="50000" r="50000" b="50000"/></a:path>'
    else:
        shade = f'<a:lin ang="{int(angle * 60000) % 21600000}" scaled="0"/>'
    return f'<a:gradFill xmlns:a="{A_NS}" rotWithShape="1"><a:gsLst>{gs}</a:gsLst>{shade}</a:gradFill>'


def _replace_fill(spPr, xml: str):
    for tag in ("a:noFill", "a:solidFill", "a:gradFill", "a:blipFill", "a:pattFill", "a:grpFill"):
        for el in spPr.findall(qn(tag)):
            spPr.remove(el)
    new = etree.fromstring(xml)
    geom = spPr.find(qn("a:prstGeom"))
    if geom is None:
        geom = spPr.find(qn("a:custGeom"))
    if geom is not None:
        geom.addnext(new)
    else:
        xfrm = spPr.find(qn("a:xfrm"))
        (xfrm.addnext(new) if xfrm is not None else spPr.insert(0, new))


def fill(shape, color=None, alpha=None, gradient=None, angle=90, radial=False, galphas=None, gpos=None):
    spPr = shape._element.spPr
    if gradient:
        _replace_fill(spPr, grad_xml(gradient, angle, radial, galphas, gpos))
    elif color:
        _replace_fill(spPr, f'<a:solidFill xmlns:a="{A_NS}">{_clr(color, alpha)}</a:solidFill>')
    else:
        _replace_fill(spPr, f'<a:noFill xmlns:a="{A_NS}"/>')


def line(shape, color=None, width_pt=1.0, alpha=None, dash=None):
    ln = shape.line
    if not color:
        ln.fill.background()
        return
    ln.width = Pt(width_pt)
    spPr = shape._element.spPr
    lnel = spPr.find(qn("a:ln"))
    for el in list(lnel):
        if el.tag in (qn("a:noFill"), qn("a:solidFill"), qn("a:gradFill"), qn("a:prstDash")):
            lnel.remove(el)
    lnel.insert(0, etree.fromstring(f'<a:solidFill xmlns:a="{A_NS}">{_clr(color, alpha)}</a:solidFill>'))
    if dash:
        lnel.append(etree.fromstring(f'<a:prstDash xmlns:a="{A_NS}" val="{dash}"/>'))


def effects(shape, shadow: str | None = None, blur_pt=14, dist_pt=4, alpha=0.14, glow: str | None = None,
            glow_pt=12, glow_alpha=0.5, soft_edge_pt=None):
    """Đặt effectLst tường minh (mặc định: không hiệu ứng, tránh bóng thừa từ theme)."""
    spPr = shape._element.spPr
    for el in spPr.findall(qn("a:effectLst")):
        spPr.remove(el)
    inner = ""
    if glow:
        inner += f'<a:glow rad="{int(glow_pt * 12700)}">{_clr(glow, glow_alpha)}</a:glow>'
    if shadow:
        inner += (f'<a:outerShdw blurRad="{int(blur_pt * 12700)}" dist="{int(dist_pt * 12700)}" dir="5400000" '
                  f'algn="t" rotWithShape="0">{_clr(shadow, alpha)}</a:outerShdw>')
    if soft_edge_pt:
        inner += f'<a:softEdge rad="{int(soft_edge_pt * 12700)}"/>'
    el = etree.fromstring(f'<a:effectLst xmlns:a="{A_NS}">{inner}</a:effectLst>')
    ln = spPr.find(qn("a:ln"))
    if ln is not None:
        ln.addnext(el)
    else:
        spPr.append(el)


# ---------- hình khối ----------

_KINDS = {"rect": MSO_SHAPE.RECTANGLE, "round": MSO_SHAPE.ROUNDED_RECTANGLE, "ellipse": MSO_SHAPE.OVAL,
          "chevron": MSO_SHAPE.CHEVRON, "pentagon": MSO_SHAPE.PENTAGON, "arrow": MSO_SHAPE.RIGHT_ARROW,
          "diamond": MSO_SHAPE.DIAMOND, "triangle": MSO_SHAPE.ISOSCELES_TRIANGLE, "arc": MSO_SHAPE.ARC,
          "tear": MSO_SHAPE.TEAR, "wave": MSO_SHAPE.WAVE, "donut": MSO_SHAPE.DONUT}


def box(slide, x, y, w, h, color=None, kind="rect", radius=0.12, alpha=None, gradient=None, angle=90,
        radial=False, galphas=None, line_color=None, line_w=1.0, line_alpha=None, shadow=None, name=None,
        rotation=0.0, glow=None, glow_pt=12, glow_alpha=0.5, soft_edge=None, dash=None):
    shp = slide.shapes.add_shape(_KINDS[kind], E(x), E(y), E(max(w, 0.01)), E(max(h, 0.01)))
    if kind == "round":
        shp.adjustments[0] = max(0.0, min(0.5, radius))
    fill(shp, color, alpha, gradient, angle, radial, galphas)
    line(shp, line_color, line_w, line_alpha, dash)
    effects(shp, shadow=shadow, glow=glow, glow_pt=glow_pt, glow_alpha=glow_alpha, soft_edge_pt=soft_edge)
    if rotation:
        shp.rotation = rotation
    if name:
        shp.name = name
    shp.text_frame.text = ""
    return shp


def poly(slide, pts, color=None, alpha=None, gradient=None, angle=90, name=None, line_color=None, line_w=1.0,
         closed=True, shadow=None):
    """Đa giác/đường tự do từ danh sách điểm (đơn vị inch)."""
    scale = 1000
    fb = slide.shapes.build_freeform(int(pts[0][0] * scale), int(pts[0][1] * scale), scale=EMU / scale)
    fb.add_line_segments([(int(px * scale), int(py * scale)) for px, py in pts[1:]], close=closed)
    shp = fb.convert_to_shape()
    if closed:
        fill(shp, color, alpha, gradient, angle)
    else:
        fill(shp, None)
    line(shp, line_color, line_w)
    effects(shp, shadow=shadow)
    if name:
        shp.name = name
    return shp


def paths(slide, rings, color=None, alpha=None, closed=True, line_color=None, line_w=1.0, line_alpha=None,
          name=None):
    """Một hình gồm nhiều nét (custGeom nhiều <a:path>): rings là một đường [(x,y),…] hoặc list các đường (inch).
    Gộp nhiều nét thành một hình giúp slide nhẹ và họa tiết sửa/di chuyển được như một khối trong PowerPoint."""
    if not isinstance(rings[0][0], (list, tuple)):
        rings = [rings]
    xs = [p[0] for r in rings for p in r]
    ys = [p[1] for r in rings for p in r]
    x0, y0 = min(xs), min(ys)
    w, h = max(max(xs) - x0, 0.01), max(max(ys) - y0, 0.01)
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, E(x0), E(y0), E(w), E(h))
    P = lambda p: f'<a:pt x="{int(round((p[0] - x0) * EMU))}" y="{int(round((p[1] - y0) * EMU))}"/>'
    body = "".join(
        f'<a:path w="{int(E(w))}" h="{int(E(h))}"><a:moveTo>{P(r[0])}</a:moveTo>'
        + "".join(f"<a:lnTo>{P(p)}</a:lnTo>" for p in r[1:]) + ("<a:close/>" if closed else "") + "</a:path>"
        for r in rings)
    geo = etree.fromstring(f'<a:custGeom xmlns:a="{A_NS}"><a:avLst/><a:gdLst/><a:ahLst/><a:cxnLst/>'
                           f'<a:rect l="0" t="0" r="r" b="b"/><a:pathLst>{body}</a:pathLst></a:custGeom>')
    spPr = shp._element.spPr
    old = spPr.find(qn("a:prstGeom"))
    old.addprevious(geo)
    spPr.remove(old)
    fill(shp, color if closed else None, alpha)
    line(shp, line_color, line_w, line_alpha)
    effects(shp)
    if name:
        shp.name = name
    shp.text_frame.text = ""
    return shp


# ---------- chữ ----------

_ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT, "j": PP_ALIGN.JUSTIFY}
_ANCH = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


def text(slide, x, y, w, h, content, font="Segoe UI", size=18, color="222222", bold=False, italic=False,
         align="l", anchor="t", spacing=1.15, after=0, caps=False, name=None, bullet=None, bullet_color=None,
         indent=0.0, margin=0.0, letter=0.0, autofit=False, shape=None):
    """content: chuỗi (\\n = đoạn mới) hoặc list đoạn; mỗi đoạn là chuỗi hoặc list (chuỗi, {kiểu run})."""
    shp = shape or slide.shapes.add_textbox(E(x), E(y), E(w), E(h))
    tf = shp.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, E(margin))
    tf.vertical_anchor = _ANCH[anchor]
    paras = content.split("\n") if isinstance(content, str) else list(content)
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = _ALIGN[align]
        if after:
            p.space_after = Pt(after)
        runs = [(para, {})] if isinstance(para, str) else para
        # giãn dòng cố định theo điểm => chiều cao đo trước khớp đúng khi hiển thị, với mọi phông
        p.line_spacing = Pt(round(max(st.get("size", size) for _, st in runs) * spacing, 1))
        for rt, st in runs:
            r = p.add_run()
            r.text = rt.upper() if (st.get("caps", caps)) else rt
            f = r.font
            f.name = st.get("font", font)
            f.size = Pt(st.get("size", size))
            f.bold = st.get("bold", bold)
            f.italic = st.get("italic", italic)
            f.color.rgb = rgb(st.get("color", color))
            if letter:
                r._r.get_or_add_rPr().set("spc", str(int(letter * 100)))
        if bullet:
            _bullet(p, bullet, bullet_color or color, indent)
    bp = tf._txBody.find(qn("a:bodyPr"))
    if autofit:
        for el in list(bp):
            bp.remove(el)
        bp.append(etree.fromstring(f'<a:normAutofit xmlns:a="{A_NS}"/>'))
    if name:
        shp.name = name
    return shp


def _bullet(p, char, color, indent_in):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(indent_in * EMU)))
    pPr.set("indent", str(-int(indent_in * EMU)))
    for tag in ("a:buClr", "a:buFont", "a:buChar", "a:buNone"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    pPr.append(etree.fromstring(f'<a:buClr xmlns:a="{A_NS}">{_clr(color)}</a:buClr>'))
    # ★ ❀ ✿ ✦ ❖ ▶ … (bullet theo họa tiết) chỉ có trong Segoe UI Symbol
    face = "Segoe UI" if char in "●■•–—♦" else "Segoe UI Symbol"
    pPr.append(etree.fromstring(f'<a:buFont xmlns:a="{A_NS}" typeface="{face}"/>'))
    pPr.append(etree.fromstring(f'<a:buChar xmlns:a="{A_NS}" char="{char}"/>'))


# ---------- ảnh ----------

def picture(slide, img, x, y, w, h, mode="cover", radius=0.0, name=None, line_color=None, shadow=None):
    """Đặt ảnh vào khung: cover = cắt cho đầy khung, contain = thu nhỏ vừa khung (căn giữa)."""
    stream = io.BytesIO(img.data)
    r_img, r_box = img.ratio, w / h
    if mode == "contain":
        if r_img > r_box:
            nw, nh = w, w / r_img
        else:
            nw, nh = h * r_img, h
        pic = slide.shapes.add_picture(stream, E(x + (w - nw) / 2), E(y + (h - nh) / 2), E(nw), E(nh))
    else:
        pic = slide.shapes.add_picture(stream, E(x), E(y), E(w), E(h))
        if r_img > r_box:  # ảnh rộng hơn khung: cắt hai bên
            c = (1 - r_box / r_img) / 2
            pic.crop_left = pic.crop_right = c
        else:
            c = (1 - r_img / r_box) / 2
            pic.crop_top = c * 0.6
            pic.crop_bottom = c * 1.4
    if radius > 0:
        geom = pic._element.spPr.find(qn("a:prstGeom"))
        geom.set("prst", "roundRect")
        av = geom.find(qn("a:avLst"))
        if av is None:
            av = etree.SubElement(geom, qn("a:avLst"))
        for g in list(av):
            av.remove(g)
        av.append(etree.fromstring(f'<a:gd xmlns:a="{A_NS}" name="adj" fmla="val {int(radius * 100000)}"/>'))
    if line_color:
        pic.line.color.rgb = rgb(line_color)
        pic.line.width = Pt(1)
    if shadow:
        effects(pic, shadow=shadow)
    if name:
        pic.name = name
    return pic


# ---------- nền slide ----------

def background(slide, colors, angle=135, radial=False):
    bg = slide.background.fill
    if len(colors) == 1:
        bg.solid()
        bg.fore_color.rgb = rgb(colors[0])
        return
    bg.gradient()
    bgPr = slide._element.find(qn("p:cSld")).find(qn("p:bg")).find(qn("p:bgPr"))
    old = bgPr.find(qn("a:gradFill"))
    new = etree.fromstring(grad_xml(colors, angle, radial))
    bgPr.replace(old, new)
